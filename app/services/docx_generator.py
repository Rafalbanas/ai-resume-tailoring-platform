from pathlib import Path

from app.models.candidate import CandidateProfile
from app.models.resume import TailoredResume
from app.services.fact_catalog import FactCatalog
from app.services.languages import profile_languages
from app.services.professional_summary import deduplicate_summary
from app.services.resume_export import contacts, link_display, prepare_resume


def add_link(paragraph, text: str, target: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def _generate_docx_local(
    resume: TailoredResume,
    profile: CandidateProfile,
    output: Path,
    *,
    photo_path: Path | None = None,
    photo_enabled: bool = False,
    template_name: str = "modern_sidebar",
) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    if not profile.personal.name.strip():
        raise ValueError("Candidate name is missing from the active master profile.")
    resume = prepare_resume(resume, profile)
    resume.professional_summary = deduplicate_summary(resume.professional_summary, FactCatalog(profile))
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.65)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 2"].font.size = Pt(10.5)
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    if photo_enabled and template_name == "modern_sidebar" and photo_path and photo_path.is_file():
        actual_photo = photo_path
        if actual_photo.suffix.casefold() == ".webp":
            sibling_jpg = actual_photo.with_name("current.jpg")
            if sibling_jpg.is_file():
                actual_photo = sibling_jpg
            else:
                import tempfile

                from PIL import Image
                with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp:
                    tmp_path = Path(tmp.name)
                try:
                    with Image.open(actual_photo) as img:
                        img.convert("RGB").save(tmp_path, "JPEG")
                    actual_photo = tmp_path
                except Exception:
                    actual_photo = None
        if actual_photo and actual_photo.is_file():
            photo = document.add_paragraph()
            photo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            photo.add_run().add_picture(str(actual_photo), width=Inches(1.15))
    title = document.add_heading(profile.personal.name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph()
    for index, contact in enumerate(contacts(profile)):
        if index:
            paragraph.add_run(" · ")
        if contact.target:
            add_link(paragraph, contact.display, contact.target)
        else:
            paragraph.add_run(contact.display)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(resume.headline).alignment = WD_ALIGN_PARAGRAPH.CENTER
    if resume.professional_summary:
        document.add_heading("Professional Summary", level=1)
        document.add_paragraph(resume.professional_summary)
    if resume.core_skills:
        document.add_heading("Core Skills", level=1)
        document.add_paragraph(" • ".join(resume.core_skills))
    if resume.currently_learning:
        document.add_heading("Currently learning", level=1)
        for item in resume.currently_learning:
            document.add_paragraph(f"{item.name} — {item.status}", style="List Bullet")
    if profile_languages(profile):
        document.add_heading("Languages", level=1)
        document.add_paragraph(" · ".join(profile_languages(profile)))
    if resume.experience:
        document.add_heading("Experience", level=1)
        for item in resume.experience:
            document.add_heading(f"{item.title} — {item.company}", level=2)
            date_paragraph = document.add_paragraph(item.dates)
            date_paragraph.paragraph_format.keep_with_next = True
            for bullet in item.bullets:
                document.add_paragraph(bullet.text, style="List Bullet")
    if resume.education:
        document.add_heading("Education", level=1)
        for item in resume.education:
            education_paragraphs = [document.add_paragraph(f"{item.qualification} — {item.institution} ({item.dates})")]
            if getattr(item, "specialisation", None):
                education_paragraphs.append(document.add_paragraph(item.specialisation))
            if getattr(item, "thesis_subline", None):
                education_paragraphs.append(document.add_paragraph(item.thesis_subline))
            for index, entry in enumerate(education_paragraphs):
                entry.paragraph_format.keep_together = True
                entry.paragraph_format.keep_with_next = index < len(education_paragraphs) - 1
    if resume.projects:
        document.add_heading("Selected Projects", level=1)
        for item in resume.projects:
            document.add_heading(item.name, level=2)
            if item.url:
                link_paragraph = document.add_paragraph()
                link_paragraph.paragraph_format.keep_with_next = True
                add_link(link_paragraph, link_display(item.url), item.url)
            paragraphs = [document.add_paragraph(item.description)]
            if item.technologies:
                paragraphs.append(document.add_paragraph(" · ".join(item.technologies)))
            for index, paragraph in enumerate(paragraphs):
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = index < len(paragraphs) - 1
    if resume.certifications:
        document.add_heading("Certifications", level=1)
        for item in resume.certifications:
            document.add_paragraph(" · ".join(filter(None, [item.name, item.issuer, item.date])))
    if getattr(resume, "interests", None):
        document.add_heading("Interests", level=1)
        document.add_paragraph(" • ".join(resume.interests))
    # Reorder/hide existing section XML without rewriting verified content.
    if resume.section_order or resume.excluded_sections:
        names = {'Professional Summary': 'summary', 'Core Skills': 'skills', 'Currently learning': 'learning',
                 'Languages': 'languages', 'Experience': 'experience', 'Education': 'education',
                 'Selected Projects': 'projects', 'Certifications': 'certifications', 'Interests': 'interests'}
        groups, active = {}, None
        for paragraph in document.paragraphs:
            if paragraph.style.name == 'Heading 1':
                active = names.get(paragraph.text)
            if active:
                groups.setdefault(active, []).append(paragraph._p)
        body = document._element.body
        for elements in groups.values():
            for element in elements:
                body.remove(element)
        for key in dict.fromkeys([*resume.section_order, *groups]):
            if key not in resume.excluded_sections:
                for element in groups.get(key, []):
                    body.insert(len(body) - 1, element)
    section.footer_distance = Inches(0.25)
    if document.paragraphs:
        document.paragraphs[-1].paragraph_format.keep_with_next = True
    consent = document.add_paragraph()
    consent.paragraph_format.keep_together = True
    consent.text = ("I consent to the processing of my personal data for recruitment purposes in accordance with applicable data protection law (GDPR).")
    for run in consent.runs:
        run.font.size = Pt(9)
    document.save(output)


def generate_docx(resume, profile, output, *, photo_path=None, photo_enabled=None, template_name="modern_sidebar"):
    from app.services.cv_content_policy import check_export_content
    check_export_content(resume)
    contacts(profile)
    from app.services.processing import decoded, encoded, process
    if photo_path and photo_path.with_name('current.jpg').is_file():
        photo_path = photo_path.with_name('current.jpg')
    photo = encoded(photo_path.read_bytes()) if photo_path and photo_path.is_file() and photo_path.stat().st_size <= 10_000_000 else ''
    result = process('docx', {'resume': resume.model_dump(mode='json'), 'profile': profile.model_dump(mode='json'),
        'photo': photo, 'options': {'photo_enabled': photo_enabled, 'template_name': template_name}})
    output.write_bytes(decoded(result))
