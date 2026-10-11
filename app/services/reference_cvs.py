from __future__ import annotations

import hashlib
import io
import json
import os
import re
import statistics
import zipfile
from collections import Counter
from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader

ALLOWED_SUFFIXES = {".pdf", ".docx"}
SECTION_ALIASES = {
    "summary": {"summary", "professional summary", "profile", "professional profile", "about me"},
    "skills": {"skills", "core skills", "technical skills", "key skills", "competencies"},
    "experience": {"experience", "work experience", "professional experience", "employment history"},
    "education": {"education", "academic background"},
    "projects": {"projects", "selected projects", "personal projects"},
    "certifications": {"certifications", "certificates"},
}
DEFAULT_SECTION_ORDER = ["summary", "experience", "education", "projects", "certifications"]


class ReferenceCVError(ValueError):
    pass


def _safe_upload_name(filename: str) -> str:
    name = Path(filename).name
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", Path(name).stem).strip(".-")[:100] or "reference-cv"
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ReferenceCVError("Only PDF and DOCX reference CVs are supported.")
    return f"{stem}{suffix}"


def _tokens(value: str) -> set[str]:
    return set(re.findall(r"[a-z0-9+#.]{2,}", value.casefold()))


class ReferenceCVLibrary:
    def __init__(self, directory: Path, max_bytes: int = 10_000_000):
        self.directory = directory
        self.index_path = directory / "index.json"
        self.max_bytes = max_bytes
        self.directory.mkdir(parents=True, exist_ok=True, mode=0o700)
        try:
            self.directory.chmod(0o700)
        except OSError:
            pass
        self._items: list[dict] = []

    def refresh(self) -> None:
        existing = self._load_index()
        by_filename = {item.get("stored_filename"): item for item in existing}
        items = []
        for path in sorted(self.directory.iterdir()):
            if not path.is_file() or path.name == "index.json" or path.suffix.lower() not in ALLOWED_SUFFIXES:
                continue
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            cached = by_filename.get(path.name)
            if cached and cached.get("sha256") == checksum:
                items.append(cached)
                continue
            items.append(self._ingest(path, checksum))
        self._items = items
        self._write_index()

    def list(self) -> list[dict]:
        return [dict(item) for item in sorted(self._items, key=lambda item: item["original_filename"].casefold())]

    def upload(self, filename: str, content: bytes) -> dict:
        safe_name = _safe_upload_name(filename)
        if not content:
            raise ReferenceCVError("The uploaded CV is empty.")
        if len(content) > self.max_bytes:
            raise ReferenceCVError("The uploaded CV exceeds the size limit.")
        if safe_name.endswith(".docx"):
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    expanded = sum(member.file_size for member in archive.infolist())
                    if expanded > min(40_000_000, self.max_bytes * 10) or len(archive.infolist()) > 2000:
                        raise ReferenceCVError("DOCX expanded content exceeds the processing limit.")
                    if "word/document.xml" not in archive.namelist():
                        raise ReferenceCVError("Upload is not a Word document.")
            except zipfile.BadZipFile as exc:
                raise ReferenceCVError("Upload is not a valid DOCX archive.") from exc
        elif not content.lstrip().startswith(b"%PDF-"):
            raise ReferenceCVError("Upload is not a valid PDF document.")
        stored_name = f"{uuid4().hex}_{safe_name}"
        path = self.directory / stored_name
        path.write_bytes(content)
        try:
            path.chmod(0o600)
        except OSError:
            pass
        record = self._ingest(path, hashlib.sha256(content).hexdigest(), original_filename=filename)
        self._items.append(record)
        self._write_index()
        return dict(record)

    def remove(self, reference_id: str) -> None:
        item = next((value for value in self._items if value["id"] == reference_id), None)
        if not item:
            raise FileNotFoundError(reference_id)
        path = self.directory / item["stored_filename"]
        if path.parent.resolve() != self.directory.resolve():
            raise FileNotFoundError(reference_id)
        path.unlink(missing_ok=True)
        self._items = [value for value in self._items if value["id"] != reference_id]
        self._write_index()

    def select(self, role: str, description: str, limit: int = 2) -> list[dict]:
        query = _tokens(f"{role} {description}")

        def score(item: dict) -> float:
            reference = _tokens(
                " ".join(
                    [item.get("target_role", ""), item.get("summary", "")]
                    + item.get("skills", [])
                    + item.get("experience_bullets", [])
                )
            )
            return len(query & reference) / max(1, len(query | reference))

        selected = sorted(
            (item for item in self._items if item.get("status") == "ready"),
            key=lambda item: (score(item), item.get("target_role", "")),
            reverse=True,
        )[:limit]
        return [
            {
                "target_role": item["target_role"],
                "summary_example": item["summary"],
                "skills_example": item["skills"],
                "experience_bullet_examples": item["experience_bullets"],
                "section_order": item["section_order"],
            }
            for item in selected
        ]

    def layout_guide(self) -> dict:
        ready = [item for item in self._items if item.get("status") == "ready"]
        if not ready:
            return {
                "summary_max_chars": 520,
                "max_skills": 12,
                "max_bullets_per_role": 4,
                "section_order": DEFAULT_SECTION_ORDER,
            }
        summary_lengths = [len(item["summary"]) for item in ready if item.get("summary")]
        skill_counts = [len(item["skills"]) for item in ready if item.get("skills")]
        bullet_counts = [len(item["experience_bullets"]) for item in ready if item.get("experience_bullets")]
        orders = Counter(tuple(item["section_order"]) for item in ready if item.get("section_order"))
        order = list(orders.most_common(1)[0][0]) if orders else DEFAULT_SECTION_ORDER
        order = ["summary"] + [section for section in order if section in DEFAULT_SECTION_ORDER and section != "summary"]
        order.extend(section for section in DEFAULT_SECTION_ORDER if section not in order)
        return {
            "summary_max_chars": min(700, max(280, int(statistics.median(summary_lengths or [520])))),
            "max_skills": min(16, max(8, int(statistics.median(skill_counts or [12])))),
            "max_bullets_per_role": min(4, max(2, round(statistics.median(bullet_counts or [4]) / 2))),
            "section_order": order,
        }

    def _load_index(self) -> list[dict]:
        try:
            payload = json.loads(self.index_path.read_text(encoding="utf-8"))
            return payload.get("items", []) if isinstance(payload, dict) else []
        except (OSError, json.JSONDecodeError):
            return []

    def _write_index(self) -> None:
        temp = self.directory / f".index-{uuid4().hex}.tmp"
        temp.write_text(json.dumps({"version": 1, "items": self._items}, indent=2, ensure_ascii=False), encoding="utf-8")
        try:
            temp.chmod(0o600)
        except OSError:
            pass
        os.replace(temp, self.index_path)

    def _ingest(self, path: Path, checksum: str, original_filename: str | None = None) -> dict:
        record = {
            "id": hashlib.sha256(path.name.encode()).hexdigest()[:16],
            "original_filename": original_filename or re.sub(r"^[a-f0-9]{32}_", "", path.name),
            "stored_filename": path.name,
            "sha256": checksum,
            "status": "ready",
            "target_role": "",
            "summary": "",
            "skills": [],
            "experience_bullets": [],
            "section_order": [],
            "error": "",
        }
        try:
            text = self._extract_text(path)
            parsed = self._parse(text, record["original_filename"])
            record.update(parsed)
            if not text.strip():
                raise ReferenceCVError("No selectable text was found in the document.")
        except Exception as exc:
            record["status"] = "error"
            record["error"] = str(exc)[:240]
        return record

    @staticmethod
    def _extract_text(path: Path) -> str:
        if path.suffix.lower() == ".pdf":
            reader = PdfReader(path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(paragraphs + cells)

    @staticmethod
    def _parse(text: str, filename: str) -> dict:
        lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
        lines = [line for line in lines if line]
        alias_lookup = {alias: section for section, aliases in SECTION_ALIASES.items() for alias in aliases}
        sections: dict[str, list[str]] = {name: [] for name in SECTION_ALIASES}
        order = []
        current = ""
        preamble = []
        for line in lines:
            normalized = re.sub(r"[^a-z ]", "", line.casefold()).strip()
            heading = alias_lookup.get(normalized)
            if heading:
                current = heading
                if heading not in order:
                    order.append(heading)
                continue
            if current:
                sections[current].append(line)
            else:
                preamble.append(line)
        role_candidates = [
            line
            for line in preamble[:15]
            if 2 <= len(line.split()) <= 9
            and "@" not in line
            and not re.search(r"https?://|linkedin|github|\+?\d[\d ()-]{7}", line, re.I)
        ]
        filename_role = re.sub(r"[_-]+", " ", Path(filename).stem)
        filename_role = re.sub(r"\b(cv|resume|final|updated|template|\d{1,4})\b", " ", filename_role, flags=re.I)
        target_role = role_candidates[-1] if role_candidates else re.sub(r"\s+", " ", filename_role).strip()
        summary = " ".join(sections["summary"])[:1600]
        if not summary:
            summary = " ".join(preamble[-3:])[:800]
        skills_text = " | ".join(sections["skills"])
        skills = [value.strip(" •·-") for value in re.split(r"[,|•·;\n]", skills_text) if value.strip(" •·-")]
        bullets = []
        for line in sections["experience"]:
            clean = line.lstrip("•·-– ").strip()
            if line[:1] in "•·-–" or len(clean.split()) >= 8:
                bullets.append(clean[:500])
        return {
            "target_role": target_role[:180],
            "summary": summary,
            "skills": skills[:30],
            "experience_bullets": bullets[:20],
            "section_order": order or DEFAULT_SECTION_ORDER,
        }
