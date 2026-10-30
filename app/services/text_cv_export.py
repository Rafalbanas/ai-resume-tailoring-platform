"""Content-preserving text CV exports, executed only in the document sandbox."""

import base64
from html import escape
from io import BytesIO

from app.services.cv_editor import split_sections, validate_text


def section_parts(text):
    for section in split_sections(text):
        lines = section['text'].splitlines(keepends=True)
        heading = section['title'] != 'Header / unclassified content'
        yield (lines[0].strip(), ''.join(lines[1:])) if heading else ('', section['text'])


def export_to_pdf(text: str, template: str) -> bytes:
    from weasyprint import HTML

    parts = []
    for title, content in section_parts(text):
        if title:
            parts.append(f'<h2>{escape(title)}</h2>')
        if content:
            parts.append(f'<div class="text">{escape(content)}</div>')
    modern = 'h2 { color: #18384b; border-bottom: 1px solid #b0c6cf; padding-bottom: 3pt; }' if template == 'modern_sidebar' else ''
    markup = '''<!doctype html><html lang="en"><head><meta charset="utf-8"><style>
        @page { size: A4; margin: 14mm; }
        body { font-family: Arial, Carlito, sans-serif; font-size: 10.5pt; line-height: 1.25;
               margin: 0; overflow-wrap: anywhere; }
        .text { white-space: pre-wrap; }
        h2 { font-size: 12pt; margin: 10pt 0 4pt; break-after: avoid; }
    ''' + modern + '</style></head><body>' + ''.join(parts) + '</body></html>'
    # Escaping prevents text from introducing URLs/resources. Deny fetching as a second boundary.
    def no_resources(url, **kwargs):
        raise ValueError('External resources are not supported in text CV exports.')
    return HTML(string=markup, url_fetcher=no_resources).write_pdf()


def export_to_docx(text: str, template: str) -> bytes:
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Mm(210), Mm(297)
    page.top_margin = page.bottom_margin = page.left_margin = page.right_margin = Mm(14)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size, normal.font.bold = 'Arial', Pt(10.5), False
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15
    heading = document.styles['Heading 1']
    heading.font.name, heading.font.size = 'Arial', Pt(12)
    if template == 'modern_sidebar':
        heading.font.color.rgb = RGBColor(24, 56, 75)
    for title, content in section_parts(text):
        if title:
            document.add_heading(title, level=1)
        # Preserve line breaks without turning an entire section into an unbreakable paragraph.
        for line in content.splitlines():
            document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def text_cv_export_op(payload: dict) -> str:
    text = validate_text(payload.get('text'))
    kind, template = payload.get('kind'), payload.get('template', 'ats_classic')
    if kind not in {'pdf', 'docx'} or template not in {'ats_classic', 'modern_sidebar'}:
        raise ValueError('Choose PDF/DOCX and an existing CV template.')
    content = export_to_pdf(text, template) if kind == 'pdf' else export_to_docx(text, template)
    return base64.b64encode(content).decode('ascii')
