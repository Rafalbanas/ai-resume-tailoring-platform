from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.models.candidate import Education
from app.models.job import JobAnalysis, JobRequest, JobRequirement, RequirementComponent
from app.models.resume import TailoredResume
from app.models.skills import SkillEvidence, SkillsDocument, VerifiedSkill
from app.services.analysis_validator import AnalysisValidator
from app.services.docx_generator import generate_docx
from app.services.evidence_matching import EvidenceMatcher
from app.services.evidence_strength import evidence_strength
from app.services.fact_catalog import FactCatalog
from app.services.fact_validator import FactValidator
from app.services.pdf_generator import PDFGenerator
from app.services.profile_enrichment import apply_enrichment
from app.services.skills_bank import SkillsBank


@pytest.fixture
def academic(profile, tmp_path):
    fact = "Used PyTorch and NumPy at a basic level in academic coursework during MSc Computer Science studies."
    profile.education = [
        Education(
            institution="Example University",
            qualification="MSc Computer Science",
            specialisation="Intelligent Systems Programming",
            facts=[fact],
        )
    ]
    profile.skills["python_ml_data_processing"] = ["Basic PyTorch (academic)", "Basic NumPy (academic)"]
    entry = next(e for e in FactCatalog(profile).prompt_entries if e.kind == "education" and e.text == fact)
    skills = [
        VerifiedSkill(
            id="skill_" + name.casefold(),
            name=name,
            category="Python / ML",
            level="basic",
            verified=True,
            allowed_in_cv=True,
            cv_wording=f"Basic {name} (academic)",
            evidence=[
                SkillEvidence(source_type="education", source_id=entry.source_id, description="Confirmed academic use.")
            ],
        )
        for name in ["PyTorch", "NumPy"]
    ]
    path = tmp_path / "skills.json"
    path.write_text(SkillsDocument(skills=skills).model_dump_json())
    return profile, SkillsBank(path, profile), entry


@pytest.mark.parametrize(
    "name", ["Basic knowledge of PyTorch", "NumPy", "Academic experience with PyTorch", "Basic knowledge of torch"]
)
def test_basic_academic_knowledge_is_supported(academic, name):
    profile, bank, entry = academic
    result = EvidenceMatcher(FactCatalog(profile), bank).match(name)
    assert result.status == "strong" and entry.source_id in [e.source_id for e in result.entries]
    assert evidence_strength(entry) == "academic_use"
    assert all(bank.wording(s).startswith("Basic ") for s in bank.skills)


@pytest.mark.parametrize(
    "name",
    [
        "Commercial experience with PyTorch",
        "Professional experience with torch",
        "Professional experience with NumPy",
        "Advanced PyTorch",
        "Expert NumPy",
        "Experience with PyTorch",
        "Training neural network models with PyTorch",
        "Advanced training of AI/ML models",
        "Production fine-tuning of models with PyTorch",
        "5 years of professional NumPy experience",
    ],
)
def test_academic_library_use_cannot_establish_broader_scope(academic, name):
    profile, bank, _ = academic
    result = EvidenceMatcher(FactCatalog(profile), bank).match(name)
    assert result.status in {"partial", "missing"} and result.observed_years is None


@pytest.mark.parametrize("components", [False, True])
def test_model_strong_claim_is_capped_even_with_all_library_components(academic, components):
    profile, bank, entry = academic
    quote = "Advanced commercial experience training neural network models with PyTorch and NumPy"
    requirement = JobRequirement(
        name=quote,
        source_quote=quote,
        status="strong",
        priority="mandatory",
        reason="Overstated model suggestion.",
        source_fact_ids=[entry.source_id],
        operator="all",
        components=[
            RequirementComponent(name=name, status="strong", source_fact_ids=[entry.source_id])
            for name in ["PyTorch", "NumPy"]
        ]
        if components
        else [],
    )
    analysis = JobAnalysis(
        company="Example",
        role="ML engineer",
        strong_matches=[],
        partial_matches=[],
        missing_requirements=[],
        recommendation="APPLY",
        match_level="HIGH",
        requirements=[requirement],
    )
    result = AnalysisValidator(profile, bank).validate(analysis, "Requirements\n" + quote, semantic=True)
    assert result.requirements and all(r.status != "strong" for r in result.requirements)
    assert result.match_level != "HIGH"
    if components:
        assert all(c.status != "strong" for r in result.requirements for c in r.components)


def test_cited_degree_cannot_be_converted_into_professional_tenure(academic):
    profile, bank, entry = academic
    quote = "3 years of professional experience with PyTorch"
    requirement = JobRequirement(
        name=quote,
        source_quote=quote,
        status="strong",
        reason="Degree dates are not employment.",
        source_fact_ids=[entry.source_id],
    )
    analysis = JobAnalysis(
        company="Example",
        role="ML engineer",
        strong_matches=[],
        partial_matches=[],
        missing_requirements=[],
        recommendation="APPLY",
        match_level="HIGH",
        requirements=[requirement],
    )
    result = AnalysisValidator(profile, bank).validate(analysis, "Requirements\n" + quote, semantic=True)
    assert result.requirements[0].status != "strong"
    assert not result.requirements[0].observed_years


@pytest.mark.parametrize("layout", ["modern_sidebar", "ats_classic"])
def test_basic_academic_wording_survives_generation_and_exports(academic, tmp_path, layout):
    profile, bank, _ = academic
    job = JobRequest(
        company="Example",
        role="Python data assistant",
        job_description="Requirements\nBasic knowledge of PyTorch and NumPy",
    )
    summary = next(e for e in FactCatalog(profile).prompt_entries if e.kind == "summary")
    resume = TailoredResume(
        headline="Support engineer",
        professional_summary=summary.text,
        summary_source_fact_ids=[summary.source_id],
        core_skills=["PyTorch", "NumPy"],
        selected_skill_ids=["skill_pytorch", "skill_numpy"],
        experience=[],
    )
    checked = FactValidator(profile, bank).validate(resume, job).resume
    assert "Basic PyTorch (academic)" in checked.core_skills and "Basic NumPy (academic)" in checked.core_skills
    pdf = PDFGenerator(Path("templates"), Path("static"), tmp_path)
    fitted, _ = pdf.fit_resume(
        checked, profile, template_name=layout, one_page=True, fact_validator=FactValidator(profile, bank)
    )
    pdf.generate(fitted, profile, tmp_path / "cv.pdf", template_name=layout)
    generate_docx(fitted, profile, tmp_path / "cv.docx", template_name=layout)
    reader = PdfReader(tmp_path / "cv.pdf")
    assert len(reader.pages) == 1
    doc = Document(tmp_path / "cv.docx")
    for text in [
        " ".join(p.extract_text() for p in reader.pages),
        " ".join(p.text for p in doc.paragraphs)
        + " "
        + " ".join(p.text for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs),
    ]:
        text = " ".join(text.split())
        assert "Basic PyTorch (academic)" in text and "Basic NumPy (academic)" in text
        assert "expert" not in text.casefold()


def test_additive_education_enrichment_is_transactional_and_keeps_other_sections(academic, tmp_path):
    profile, bank, _ = academic
    path = tmp_path / "master_profile.json"
    path.write_text(profile.model_dump_json())
    before = path.read_bytes()
    bank_before = (tmp_path / "skills.json").read_bytes()
    plan = {"education_facts": [{"owner": "Example University", "facts": ["Additional confirmed academic use."]}]}
    apply_enrichment(tmp_path, plan, dry_run=True)
    assert path.read_bytes() == before and (tmp_path / "skills.json").read_bytes() == bank_before
    backup = apply_enrichment(tmp_path, plan)
    assert backup and (backup / "master_profile.json").read_bytes() == before
    from app.models.candidate import CandidateProfile

    changed = CandidateProfile.model_validate_json(path.read_text())
    assert (
        changed.experience == profile.experience
        and changed.projects == profile.projects
        and changed.skills == profile.skills
    )
    assert changed.education[0].facts == profile.education[0].facts + ["Additional confirmed academic use."]


def test_unrelated_employment_citations_cannot_supply_ml_tenure(academic):
    profile, bank, entry = academic
    profile.experience[0].start = "2020-01"
    profile.experience[0].end = "Present"
    unrelated = next(e for e in FactCatalog(profile).prompt_entries if e.kind == "experience")
    quote = "1 year of advanced professional experience with PyTorch"
    requirement = JobRequirement(
        name=quote,
        source_quote=quote,
        status="strong",
        reason="Unsupported suggestion.",
        source_fact_ids=[entry.source_id, unrelated.source_id],
    )
    analysis = JobAnalysis(
        company="Example",
        role="ML engineer",
        strong_matches=[],
        partial_matches=[],
        missing_requirements=[],
        recommendation="APPLY",
        match_level="HIGH",
        requirements=[requirement],
    )
    result = AnalysisValidator(profile, bank).validate(analysis, "Requirements\n" + quote, semantic=True)
    assert result.requirements[0].status != "strong"
    assert not result.requirements[0].observed_years


def test_dated_relevant_employment_can_support_commercial_library_use(academic):
    profile, bank, _ = academic
    profile.experience[0].start = "2020-01"
    profile.experience[0].end = "Present"
    profile.experience[0].facts = ["Used NumPy for numerical computations in a professional application."]
    entry = next(e for e in FactCatalog(profile).prompt_entries if e.kind == "experience")
    quote = "1 year of professional experience with NumPy"
    requirement = JobRequirement(
        name=quote,
        source_quote=quote,
        status="strong",
        reason="Documented employment use.",
        source_fact_ids=[entry.source_id],
    )
    analysis = JobAnalysis(
        company="Example",
        role="Data developer",
        strong_matches=[],
        partial_matches=[],
        missing_requirements=[],
        recommendation="APPLY",
        match_level="HIGH",
        requirements=[requirement],
    )
    result = AnalysisValidator(profile, bank).validate(analysis, "Requirements\n" + quote, semantic=True)
    assert result.requirements[0].status == "strong" and result.requirements[0].observed_years >= 1


def test_basic_academic_practical_use_is_not_mistaken_for_only_theory(academic):
    profile, bank, _ = academic
    result = EvidenceMatcher(FactCatalog(profile), bank).match("Basic practical experience with NumPy")
    assert result.status == "strong" and "academic" in result.reason.casefold()


def test_library_declarations_do_not_prove_academic_practical_use(academic):
    profile, _, _ = academic
    profile.education[0].facts = ["Basic theoretical knowledge of PyTorch and NumPy during academic studies."]
    result = EvidenceMatcher(FactCatalog(profile)).match("Academic practical experience with PyTorch")
    assert result.status == "partial"


@pytest.mark.parametrize(
    "name",
    [
        "Doświadczenie komercyjne z PyTorch",
        "Ekspercka znajomość NumPy",
        "Zaawansowane trenowanie modeli AI/ML",
        "3 lata NumPy",
    ],
)
def test_polish_scope_modifiers_do_not_bypass_academic_limits(academic, name):
    profile, bank, _ = academic
    result = EvidenceMatcher(FactCatalog(profile), bank).match(name)
    assert result.status in {"partial", "missing"} and result.observed_years is None
