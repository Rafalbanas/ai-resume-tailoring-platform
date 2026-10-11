import base64
from contextlib import closing
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pypdfium2
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image, ImageChops
from pypdf import PdfReader
from weasyprint import HTML

from app.models.resume import TailoredResume
from app.services.docx_generator import generate_docx
from app.services.fact_catalog import FactCatalog
from app.services.pdf_generator import PDFGenerator
from app.services.professional_summary import build_summary, deduplicate_summary


def test_verified_repeated_clause_preserves_identity_details_and_source_ids(profile):
    identity = 'Provides Tier 2 support for Example Systems partners across Europe.'
    detail = 'Provides Tier 2 technical support to partners across Europe, resolving network and application issues.'
    profile.summary_facts = [identity]
    profile.experience[0].facts = [detail, 'Writes operational documentation in English.']
    catalog = FactCatalog(profile)
    expected = ('Provides Tier 2 technical support for Example Systems partners across Europe, '
                'resolving network and application issues.')
    assert deduplicate_summary(identity + ' ' + detail, catalog) == expected
    summary, ids = build_summary(catalog, 'Customer support')
    assert summary.startswith(expected)
    assert 'Writes operational documentation in English.' in summary
    assert len(ids) == 3
    assert deduplicate_summary(summary, catalog) == summary
    # Unknown model wording and distinct scope/numbers are never merged.
    other = 'Provides Tier 3 technical support to partners across Europe, resolving network issues.'
    profile.experience[0].facts.append(other)
    assert deduplicate_summary(identity + ' ' + other, FactCatalog(profile)) == identity + ' ' + other
    negative = 'Provides no Tier 2 support for Example Systems partners across Europe.'
    profile.summary_facts.append(negative)
    assert deduplicate_summary(negative + ' ' + detail, FactCatalog(profile)) == negative + ' ' + detail
    unknown = detail.replace('network', 'unverified security')
    assert deduplicate_summary(identity + ' ' + unknown, catalog) == identity + ' ' + unknown


def long_resume(resume_payload):
    payload = deepcopy(resume_payload)
    payload['headline'] = 'Customer Service Consultant and Technical Application Support Specialist for Enterprise Integration and Infrastructure Operations'
    payload['core_skills'] = [f'Technology {i}: application integration and operational troubleshooting' for i in range(18)]
    item = payload['experience'][0]
    item['title'] = 'Senior Enterprise Application Integration and Technical Customer Support Consultant'
    item['dates'] = 'September 2020 – September 2026'
    item['bullets'] = [{'text': f'Fact {i}: ' + 'Resolved verified customer application issues and documented the operational workflow. ' * 3,
                        'source_fact_ids': ['experience:0:fact:0']} for i in range(4)]
    payload['experience'] = [deepcopy(item) for _ in range(7)]
    payload['projects'] = [{'name': f'Project {i}: enterprise application integration and infrastructure operations',
                            'description': 'Implemented verified workflows and documented maintenance procedures. ' * 6,
                            'url': 'https://example.com/' + 'long-segment-' * 18,
                            'technologies': ['LongTechnologyIdentifier' * 3, 'Python', 'Linux'],
                            'source_fact_ids': ['project:0:description']} for i in range(3)]
    return TailoredResume.model_validate(payload)


def test_all_page_text_stays_inside_paper_and_reserved_footer(profile, resume_payload):
    generator = PDFGenerator(Path('templates'), Path('static'))
    resume = long_resume(resume_payload)
    for theme in ('modern_sidebar', 'ats_classic'):
        doc = HTML(string=generator.render_html(resume, profile, template_name=theme)).render()
        assert len(doc.pages) > 1
        for page in doc.pages:
            for box in page._page_box.descendants():
                if type(box).__name__ != 'TextBox' or not box.text.strip():
                    continue
                assert box.position_x >= -0.1, (theme, box.text)
                assert box.position_x + box.width <= page.width + 0.1, (theme, box.text)
                assert box.position_y >= -0.1, (theme, box.text)
                assert box.position_y + box.height <= page.height + 0.1, (theme, box.text)
        def check_flow(box, bottom, margin=False):
            margin = margin or type(box).__name__ == 'MarginBox'
            if type(box).__name__ == 'TextBox' and box.text.strip() and not margin:
                assert box.position_y + box.height <= bottom + .1
            for child in getattr(box, 'children', []):
                check_flow(child, bottom, margin)
        for page in doc.pages:
            check_flow(page._page_box, page._page_box.content_box_y() + page._page_box.height)
        reader = PdfReader(BytesIO(doc.write_pdf()))
        for page in reader.pages:
            text = page.extract_text()
            assert 'GDPR' in text
            assert len(text.replace('I consent to the processing of my personal data for recruitment purposes in accordance with applicable data protection law (GDPR).', '').strip()) > 50
        all_text = ' '.join(page.extract_text() for page in reader.pages)
        assert all_text.count('Fact 0:') == 7
        assert 'Project 2:' in all_text


def test_preview_pages_match_pdf_and_do_not_mutate_saved_resume(profile, resume_payload, tmp_path):
    resume = TailoredResume.model_validate(resume_payload)
    before = resume.model_dump_json()
    generator = PDFGenerator(Path('templates'), Path('static'), tmp_path)
    for theme in ('modern_sidebar', 'ats_classic'):
        pdf = tmp_path / f'{theme}.pdf'
        generator.generate(resume, profile, pdf, template_name=theme)
        preview = BeautifulSoup(generator.render_preview(resume, profile, template_name=theme,
                                                         interactive_photo=True), 'html.parser')
        assert len(preview.select('.preview-page')) == len(PdfReader(pdf).pages)
        assert all(page.select_one('.page-image')['src'].startswith('data:image/png;base64,')
                   for page in preview.select('.preview-page'))
        with closing(pypdfium2.PdfDocument(pdf)) as rendered:
            with closing(rendered[0]) as page:
                bitmap = page.render(scale=1.6)
                try:
                    expected = bitmap.to_pil().convert('RGB')
                    actual = Image.open(BytesIO(base64.b64decode(preview.select_one('.page-image')['src'].split(',')[1]))).convert('RGB')
                    assert ImageChops.difference(expected, actual).getbbox() is None
                finally:
                    bitmap.close()
        assert bool(preview.find(id='quick-photo-avatar')) == (theme == 'modern_sidebar')
        assert preview.select_one('a[href="mailto:jan@example.com"]')
        output = tmp_path / f'{theme}.docx'
        generate_docx(resume, profile, output, template_name=theme)
        doc = Document(output)
        assert 'Invented Kubernetes achievement' in '\n'.join(p.text for p in doc.paragraphs)
        assert 'GDPR' in doc.sections[0].footer.paragraphs[0].text
    assert resume.model_dump_json() == before
