from copy import deepcopy

import pytest
from docx import Document
from pypdf import PdfReader

from app.main import app
from app.models.resume import ResumeLearningSkill
from app.models.skills import SkillEvidence
from app.services.fact_validator import FactValidator
from app.services.resume_editor import apply_resume_edits
from tests.test_safe_resume_edits import form_for, saved_cv  # noqa: F401
from tests.test_skills_bank import bank, skill


def learning(name, *, course=True, enabled=True, attested=True):
    entry = skill(name, level='learning', verified=False, allowed=False, enabled=enabled)
    if attested:
        entry.evidence = [SkillEvidence(source_type='manual_verified',
            source_id=f'manual:{"learning-course" if course else "learning"}:{name.lower()}',
            description='User confirmed current learning; no completed course or practical proficiency.')]
    return entry


@pytest.fixture
def learning_cv(saved_cv, tmp_path):  # noqa: F811
    client, slug, folder, original, job = saved_cv
    skills = bank(tmp_path, app.state.profile, [learning('Ansible'), learning('Kubernetes')])
    app.state.skills_bank = skills
    job.job_description = 'Requirements: Kubernetes and Ansible for infrastructure automation.'
    validator = FactValidator(app.state.profile, skills)
    resume = validator.validate(original, job).resume
    return client, slug, folder, original, job, validator, resume


def test_relevant_attested_learning_only(learning_cv, tmp_path):
    _, _, _, original, job, _, _ = learning_cv
    skills = bank(tmp_path, app.state.profile, [learning('Ansible'), learning('Kubernetes', enabled=False),
        learning('Terraform', attested=False), learning('Rust'), learning('Python', course=False)])
    job.job_description += ' Terraform and Python.'
    checked = FactValidator(app.state.profile, skills).validate(original, job).resume
    assert [(i.name, i.status) for i in checked.currently_learning] == [
        ('Ansible', 'Course in progress'), ('Python', 'Learning in progress')]
    assert all(not skills.eligible(s) for s in skills.skills)
    assert not any('Ansible' in s or 'Kubernetes' in s for s in checked.core_skills)
    job.job_description = 'Customer service and desktop support.'
    assert not FactValidator(app.state.profile, skills).validate(original, job).resume.currently_learning


def test_model_cannot_invent_learning_or_promote_level(learning_cv):
    _, _, _, original, job, validator, _ = learning_cv
    original.currently_learning = [ResumeLearningSkill(skill_id='skill_terraform', name='Expert Terraform')]
    original.core_skills = ['Kubernetes', 'Advanced Ansible']
    checked = validator.validate(original, job)
    assert {i.name for i in checked.resume.currently_learning} == {'Ansible', 'Kubernetes'}
    assert not checked.resume.core_skills
    assert checked.warnings


def test_manual_learning_additions_validated(learning_cv):
    _, _, _, original, job, validator, resume = learning_cv
    edited = deepcopy(original)
    edited.currently_learning = resume.currently_learning
    assert not validator.validate_manual_edit(original, edited, job).issues
    edited.currently_learning[0].name = 'Expert Ansible'
    assert any('Currently learning' in i['field'] for i in validator.validate_manual_edit(original, edited, job).issues)
    edited = deepcopy(original)
    edited.core_skills.append('Basic Kubernetes')
    assert validator.validate_manual_edit(original, edited, job).issues


@pytest.mark.parametrize('template', ['ats_classic', 'modern_sidebar'])
def test_learning_export_and_removal_persist(learning_cv, template):
    client, _, _, _, job, _, resume = learning_cv
    from app.models.job import JobAnalysis
    analysis = JobAnalysis(company=job.company, role=job.role, strong_matches=[], partial_matches=[],
                           missing_requirements=[], recommendation='REVIEW', match_level='LOW')
    slug, folder = app.state.storage.save_application(job, analysis, resume, [])
    response = client.post(f'/edit/{slug}', data=form_for(client, resume, template_name=template), follow_redirects=False)
    assert response.status_code == 303, response.text
    pdf = '\n'.join(p.extract_text() for p in PdfReader(folder/'resume.pdf').pages)
    docx = '\n'.join(p.text for p in Document(folder/'resume.docx').paragraphs)
    for text in (pdf, docx):
        assert 'currently learning' in text.casefold()
        assert 'Ansible' in text and 'Kubernetes' in text and 'Course in progress' in text
    before = [p.read_bytes() for p in (app.state.settings.master_profile_path, app.state.settings.skills_path)]
    form = form_for(client, resume, currently_learning='Kubernetes — Course in progress', template_name=template)
    response = client.post(f'/edit/{slug}', data=form, follow_redirects=False)
    assert response.status_code == 303, response.text
    _, saved, _, _ = app.state.storage.load_application(slug)
    assert [i.name for i in saved.currently_learning] == ['Kubernetes']
    assert client.get(f'/edit/{slug}').status_code == 200
    assert before == [p.read_bytes() for p in (app.state.settings.master_profile_path, app.state.settings.skills_path)]
    assert 'Ansible' not in '\n'.join(p.extract_text() for p in PdfReader(folder/'resume.pdf').pages)
    assert 'Ansible' not in '\n'.join(p.text for p in Document(folder/'resume.docx').paragraphs)
    edited = apply_resume_edits(saved, {'currently_learning': ''})
    assert not edited.currently_learning
