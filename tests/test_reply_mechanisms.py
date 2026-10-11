"""Synthetic regressions for source context, evidence and honest comparisons."""
from pathlib import Path

import pytest
from docx import Document
from jinja2 import Environment, FileSystemLoader
from pypdf import PdfReader
from weasyprint import HTML

from app.models.job import JobAnalysis, JobRequest
from app.models.resume import TailoredResume
from app.models.skills import SkillEvidence, SkillsDocument, VerifiedSkill
from app.services.ai_provider import build_compact_analysis_payload, build_compact_tailor_payload
from app.services.analysis_validator import AnalysisValidator
from app.services.comparison_capture import capture_comparison, record_call
from app.services.docx_generator import generate_docx
from app.services.fact_validator import FactValidator
from app.services.pdf_generator import PDFGenerator
from app.services.skills_bank import SkillsBank


def empty():
    return JobAnalysis(company='Example', role='Customer support', strong_matches=[], partial_matches=[],
                       missing_requirements=[], recommendation='APPLY', match_level='HIGH')


def test_duties_and_stack_never_imply_prior_experience(profile):
    analysis = empty()
    analysis.missing_requirements = ['Configure and maintain customer service processes within Dynamics 365.']
    result = AnalysisValidator(profile).validate(analysis,
        'JOB DESCRIPTION\nResponsibilities\nConfigure and maintain customer service processes within Dynamics 365.\nTECH STACK\nCRM\nCustomer Support')
    assert not result.requirements
    assert len(result.responsibilities) == 1
    assert [r.name for r in result.role_competencies] == ['CRM', 'Customer Support', 'Dynamics 365']
    assert result.is_reliable and result.match_level == 'LIMITED' and result.recommendation == 'REVIEW'
    assert 'Qualification requirements are not specified.' in result.reasoning_summary


def test_semantic_duplicates_keep_quotes_but_configuration_remains_separate(profile):
    text = ('Qualifications\nKnowledge of CRM systems.\nFamiliarity with CRM.\n'
            'Knowledge of Dynamics 365.\nExperience configuring Dynamics 365.')
    result = AnalysisValidator(profile).validate(empty(), text)
    assert len(result.requirements) == 3
    assert result.requirements[0].source_quotes == ['Knowledge of CRM systems.', 'Familiarity with CRM.']
    assert result.requirements[1].status == result.requirements[2].status == 'missing'
    assert result.requirements[1].group_id != result.requirements[2].group_id


def test_evidence_semantics_uses_content_not_employer_or_title(profile):
    profile.experience[0].facts = ['Supports partners and users with application issues.',
                                 'Troubleshoots applications and resolves incidents.']
    result = AnalysisValidator(profile).validate(empty(),
        'Requirements\nExperience in customer support\nProblem solving skills\nDynamics 365\nKCS')
    assert [r.status for r in result.requirements] == ['strong', 'strong', 'missing', 'missing']
    assert all(result.match_sources[r.name] for r in result.requirements[:2])
    assert result.requirements[0].evidence == [profile.experience[0].facts[0]]
    profile.experience[0].title = 'Dynamics Customer Support Problem Solving Expert'
    profile.experience[0].facts = ['Works on unrelated assignments.']
    result = AnalysisValidator(profile).validate(empty(), 'Requirements\nCustomer Support\nProblem Solving\nDynamics 365')
    assert all(r.status == 'missing' for r in result.requirements)


def test_profile_silence_and_explicit_gap_are_distinct(profile):
    profile.summary_facts += ['No experience with Dynamics 365.']
    result = AnalysisValidator(profile).validate(empty(), 'Requirements\nDynamics 365\nKCS')
    assert [r.status for r in result.requirements] == ['absent', 'missing']


def test_skills_rank_job_over_model_order_without_padding(tmp_path, profile, resume_payload):
    skills = [VerifiedSkill(id=f'skill_{i}', name=name, category='Tools', level='hands_on', verified=True,
        allowed_in_cv=True, evidence=[SkillEvidence(source_type='manual_verified', source_id=f'manual:{i}', description='Verified fixture')])
        for i, name in enumerate(['Windows 11', 'Linux', 'Salesforce', 'Incident handling', 'Technical troubleshooting'])]
    path = tmp_path/'skills.json'
    path.write_text(SkillsDocument(skills=skills).model_dump_json())
    bank = SkillsBank(path, profile)
    candidate = TailoredResume.model_validate(resume_payload)
    candidate.core_skills = ['Windows 11', 'Linux']
    candidate.selected_skill_ids = [skills[0].id, skills[1].id]
    support_job = JobRequest(company='Example', role='Customer support', job_description='Provide customer support using CRM and incident handling.')
    result = FactValidator(profile, bank).validate(candidate, support_job)
    assert 'Hands-on Salesforce' in result.resume.core_skills
    assert not any('Windows' in s or 'Linux' in s for s in result.resume.core_skills)
    linux_job = support_job.model_copy(update={'role':'Linux operations', 'job_description':'Administer Linux servers and services.'})
    linux = FactValidator(profile, bank).validate(candidate, linux_job)
    assert linux.resume.core_skills == ['Hands-on Linux']
    assert bank.document.skills == skills


def test_summary_keeps_verified_subjectless_sentences(profile, resume_payload):
    profile.summary_facts = ['Provides customer support for partners.', 'Troubleshoots incidents and documents solutions.']
    candidate = TailoredResume.model_validate(resume_payload)
    candidate.professional_summary = 'Troubleshoots incidents and documents solutions.'
    candidate.summary_source_fact_ids = ['summary:1']
    job = JobRequest(company='Example', role='Support specialist', job_description='Provide customer support and document solutions to incidents.')
    result = FactValidator(profile).validate(candidate, job)
    assert result.resume.professional_summary == candidate.professional_summary
    assert not any('Professional summary rebuilt' in w for w in result.warnings)
    candidate.professional_summary = 'Expert Dynamics consultant.'
    rebuilt = FactValidator(profile).validate(candidate, job)
    assert 'Dynamics' not in rebuilt.resume.professional_summary
    assert 'Provides' in rebuilt.resume.professional_summary
    assert any('model tone cannot be assessed' in w for w in rebuilt.warnings)


def test_payload_has_full_job_and_experience_evidence(profile):
    job = JobRequest(company='Example', role='Support', job_description='Responsibilities\nResolve application issues and support customers.\nTECH STACK\nCRM')
    payload = build_compact_analysis_payload(job, profile)
    assert payload['job']['job_description'] == job.job_description
    assert payload['candidate']['experience'][0]['bullets'] == profile.experience[0].facts
    assert build_compact_tailor_payload(job, profile, empty())['job_description'] == job.job_description


def test_capture_is_scoped_and_records_exact_input_and_response():
    record_call('ollama', {}, 'prompt', {}, 'discarded')
    with capture_comparison() as records:
        record_call('ollama', {'stage':'analyze_job','job':'exact'}, 'prompt', {}, '{raw}')
        with capture_comparison() as nested:
            record_call('gemini', {}, 'prompt', {}, 'nested')
        assert len(records) == len(nested) == 1
    assert records[0]['raw_response'] == '{raw}'
    assert records[0]['payload']['job'] == 'exact'


def test_comparison_labels_rebuilt_summary_and_unavailable_old_output(resume_payload):
    analysis = empty()
    from app.models.job import JobRequirement
    analysis.requirements = [JobRequirement(name='CRM')]
    result = {'ok':True,'model':'test','duration_ms':1,'analysis':analysis.model_dump(),
              'resume':resume_payload, 'truth_warnings':['Professional summary rebuilt from verified source sentences']}
    html = Environment(loader=FileSystemLoader('templates')).get_template('compare.html').render(
        results={'ollama':result,'gemini':result},job={'role':'Support','company':'Example'},csrf_token='fixture',request={'state':{'authenticated':False}})
    assert 'Requirements & Evidence' in html
    assert 'does not represent the model’s tone' in html
    assert 'cannot be reconstructed' in html
    assert 'Wymagania' not in html


@pytest.mark.parametrize('theme', ['modern_sidebar', 'ats_classic'])
@pytest.mark.parametrize('long', [False, True])
def test_pdf_reserved_consent_never_forms_own_page(tmp_path, profile, resume_payload, theme, long):
    resume = TailoredResume.model_validate(resume_payload)
    if long:
        resume.experience = [resume.experience[0].model_copy(deep=True) for _ in range(25)]
    generator = PDFGenerator(Path('templates'),Path('static'),tmp_path)
    fitted, _ = generator.fit_resume(resume, profile, template_name=theme)
    assert fitted == resume
    output = tmp_path/f'{theme}.pdf'
    generator.generate(fitted,profile,output,template_name=theme)
    pages = PdfReader(output).pages
    assert (len(pages) > 1) == long
    for page in pages:
        text = page.extract_text()
        assert 'GDPR' in text and 'Example Ltd' in text
    rendered = HTML(string=generator.render_html(resume,profile,template_name=theme)).render()
    for page in rendered.pages:
        boxes = list(page._page_box.descendants())
        footer = next(b for b in boxes if type(b).__name__ == 'BlockBox' and getattr(b,'element',None) is not None and 'consent-footer' in b.element.get('class',''))
        lines = [b for b in boxes if type(b).__name__=='LineBox' and b.position_y < footer.position_y]
        assert all(b.position_y + b.height <= footer.position_y for b in lines)
        assert footer.position_y + footer.height <= page.height


def test_docx_consent_is_in_footer_and_text_is_preserved(tmp_path, profile, resume_payload):
    resume=TailoredResume.model_validate(resume_payload)
    output=tmp_path/'resume.docx'
    generate_docx(resume,profile,output)
    document=Document(output)
    assert 'GDPR' in document.sections[0].footer.paragraphs[0].text
    assert all('GDPR' not in p.text for p in document.paragraphs)
    assert resume.professional_summary in [p.text for p in document.paragraphs]


@pytest.mark.asyncio
async def test_comparison_calls_each_provider_once_and_preserves_global_retries(tmp_path, profile):
    from datetime import datetime, timezone
    from types import SimpleNamespace

    from app.services.ai_provider import GeminiQuotaError
    from app.services.mock_provider import MockAIProvider
    from app.services.task_manager import TaskKind, TaskManager, TaskRecord

    calls = []

    class QuotaProvider:
        name = 'gemini'
        model = 'fixture'
        api_key = 'fixture'
        max_retries = 3

        async def tailor(self, job, profile):
            calls.append((self.name, self.max_retries))
            raise GeminiQuotaError('Quota exceeded (HTTP 429)')

    class LocalProvider(MockAIProvider):
        name = 'ollama'

        async def tailor(self, job, profile):
            calls.append((self.name, 0))
            return await super().tailor(job, profile)

    gemini = QuotaProvider()
    state = SimpleNamespace(profile=profile, skills_bank=None, providers={'ollama':LocalProvider(), 'gemini':gemini})
    task = TaskRecord(task_id='d'*32,kind=TaskKind.COMPARE,idempotency_key='fixture',
        created_at=datetime.now(timezone.utc).isoformat(),
        payload={'job':JobRequest(company='Example',role='Support',job_description='Requirements\nCustomer Support\nProblem Solving').model_dump()})
    manager = TaskManager(tmp_path)
    manager._save_task(task)
    await manager._execute_compare(task,state)
    result = manager.get_task(task.task_id).result['results']
    assert calls == [('ollama', 0), ('gemini', 0)]
    assert gemini.max_retries == 3
    assert result['gemini']['ok'] is False and '429' in result['gemini']['error']
    assert result['ollama']['model_resume']
    assert (tmp_path/'tasks'/f'{task.task_id}.json').stat().st_mode & 0o777 == 0o600


@pytest.mark.asyncio
async def test_invalid_response_is_captured_before_schema_parsing(profile):
    import httpx

    from app.core.config import Settings
    from app.services.ai_provider import AIProviderError
    from app.services.ollama_provider import OllamaProvider

    transport = httpx.MockTransport(lambda request: httpx.Response(200,json={'message':{'content':'{"invalid":'}}))
    provider = OllamaProvider(Settings(_env_file=None),transport=transport)
    job = JobRequest(company='Example',role='Support',job_description='Requirements\nCustomer Support\nProblem Solving')
    with capture_comparison() as records:
        with pytest.raises(AIProviderError):
            await provider.analyze_job(job,profile)
    assert len(records) == 1 and records[0]['raw_response'] == '{"invalid":'
    assert records[0]['payload']['job']['job_description'] == job.job_description
