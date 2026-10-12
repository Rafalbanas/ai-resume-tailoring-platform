from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfReader

from app.models.resume import ResumeProject, TailoredResume
from app.services.docx_generator import generate_docx
from app.services.pdf_generator import PDFGenerator
from app.services.resume_export import validate_layout


def document_with_projects(resume_payload, count=3):
    r = TailoredResume.model_validate(resume_payload)
    r.professional_summary = "Verified support work with applications, networking and operational SQL queries."
    r.core_skills = ["Python", "PowerShell", "Linux", "SQL Server", "T-SQL", "Application troubleshooting"]
    r.projects = [
        {
            "name": f"Project {i}",
            "description": "Implemented a documented workflow for operational data analysis and API integration. "
            + ("Additional verified workflow detail. " * 6),
            "technologies": ["Python", "JSON", "SQL"],
            "source_fact_ids": ["fixture"],
            "url": "https://example.com/projects/" + str(i),
        }
        for i in range(count)
    ]
    r.projects = [ResumeProject.model_validate(p) for p in r.projects]
    return r


def test_project_blocks_stay_whole_and_consent_is_once(tmp_path, profile, resume_payload):
    r = document_with_projects(resume_payload)
    r.experience = []
    from app.models.resume import ResumeBullet, ResumeExperience

    for i in range(5):
        r.experience.append(
            ResumeExperience(
                company=f"Example {i}",
                title="Technical Engineer",
                dates="2021–2026",
                bullets=[
                    ResumeBullet(
                        text="Investigated application incidents, networking and integration problems using documented operational diagnostics. "
                        + f"Task {j}.",
                        source_fact_ids=["fixture"],
                    )
                    for j in range(4)
                ],
            )
        )
    g = PDFGenerator(Path("templates"), Path("static"), tmp_path)
    output = tmp_path / "two-pages.pdf"
    g.generate(r, profile, output)
    reader = PdfReader(output)
    assert len(reader.pages) >= 2
    content = [p.extract_text() for p in reader.pages]
    assert "".join(content).count("I consent to the processing") == 1
    for project in r.projects:
        page = next(text for text in content if project.name in text)
        assert project.description in page.replace("\n", " ")
        assert "Python · JSON · SQL" in page
    for exp in r.experience:
        assert exp.company in "".join(content)
    doc = g._render_document(g.render_html(r, profile))
    tail = g.page_content_metrics(doc)[-1]
    assert tail["height"] >= tail["available_height"] * 0.2
    out = tmp_path / "resume.docx"
    generate_docx(r, profile, out)
    document = Document(out)
    assert sum("I consent" in p.text for p in document.paragraphs) == 1
    for project in r.projects:
        idx = next(i for i, p in enumerate(document.paragraphs) if p.text == project.name)
        assert document.paragraphs[idx + 1].paragraph_format.keep_with_next
        assert document.paragraphs[idx + 2].paragraph_format.keep_together


def test_skills_budget_only_changes_fitted_copy(tmp_path, profile, resume_payload):
    r = document_with_projects(resume_payload, 1)
    r.core_skills = [
        f"Verified skill {i}: operational analysis and troubleshooting in a documented integration workflow"
        for i in range(18)
    ]
    before = r.model_copy(deep=True)
    g = PDFGenerator(Path("templates"), Path("static"), tmp_path)
    fit, warnings = g.fit_resume(r, profile, {})
    assert r == before
    assert len(fit.core_skills) >= 10
    assert fit.experience == before.experience and fit.education == before.education
    assert [p.name for p in fit.projects] == [p.name for p in before.projects]
    if len(fit.core_skills) < 18:
        assert any("skill" in w for w in warnings)


def test_overlap_detection_uses_actual_text_geometry():
    TextBox = type("TextBox", (), {})
    boxes = []
    for x in [10, 15]:
        b = TextBox()
        b.text = "Text"
        b.position_x = x
        b.position_y = 10
        b.width = 40
        b.height = 12
        b.children = []
        b.element = None
        boxes.append(b)
    root = SimpleNamespace(children=boxes, element=None)
    page = SimpleNamespace(width=100, height=100, _page_box=root)
    with pytest.raises(ValueError, match="overlaps"):
        validate_layout(SimpleNamespace(pages=[page]))


def test_single_page_consent_and_pdf_links(tmp_path, profile, resume_payload):
    r = TailoredResume.model_validate(resume_payload)
    path = tmp_path / "one.pdf"
    PDFGenerator(Path("templates"), Path("static"), tmp_path).generate(r, profile, path)
    reader = PdfReader(BytesIO(path.read_bytes()))
    assert len(reader.pages) == 1
    assert reader.pages[0].extract_text().count("I consent") == 1
    assert any(a.get_object().get("/A", {}).get("/URI") == "mailto:jan@example.com" for a in reader.pages[0]["/Annots"])
