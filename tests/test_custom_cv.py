from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.core.config import Settings
from app.main import app
from app.models.job import JobRequest
from app.models.resume import TailoredResume
from app.services.ai_provider import (
    ProviderResponseError,
    ProviderUnavailableError,
    build_compact_tailor_payload,
    resume_system_prompt,
)
from app.services.custom_cv import custom_analysis
from app.services.cv_editor import (
    CVEditorStore,
    EditProposal,
    EditReview,
    SectionReplacement,
    apply_proposal,
    edit_with_chat,
    split_sections,
)
from app.services.fact_validator import FactValidator
from app.services.gemini_provider import GeminiProvider
from app.services.ollama_provider import OllamaProvider, truth_rules
from app.services.resilient_provider import ResilientAIProvider
from app.services.task_manager import TaskKind, TaskManager, TaskStatus
from tests.test_recording_demo import sign_in

SOURCE = ("Ada Example\nada@example.test\n\nProfessional Summary\n"
          "I troubleshoot SQL incidents.\n\nExperience\n"
          "Example Ltd | Support Engineer | 2022–Present\nResolved customer incidents using SQL.\n\n"
          "Education\nExample University | BSc | 2021\n")


@pytest.fixture(autouse=True)
def offline_model_listing(monkeypatch):
    monkeypatch.setattr(OllamaProvider, 'available_models', AsyncMock(return_value=[]))


def custom_job(**kwargs):
    return JobRequest(company="Custom CV", role="Support CV", workflow="custom",
                      custom_instructions="Emphasize verified Python automation and technical support.", **kwargs)


def test_custom_request_needs_no_offer_but_job_mode_still_does():
    assert custom_job().job_description == ""
    with pytest.raises(ValidationError):
        JobRequest(company="Example", role="Support")
    with pytest.raises(ValidationError):
        JobRequest(company="Custom", role="CV", workflow="custom", custom_instructions="   ")


def test_custom_payload_and_truth_boundary(profile):
    job = custom_job()
    payload = build_compact_tailor_payload(job, profile, custom_analysis(job))
    assert payload["custom_instructions"] == job.custom_instructions
    assert payload["job_description"] == ""
    assert payload["workflow"] == "custom"
    assert "Never invent" in resume_system_prompt(job)
    assert "unsupported" in resume_system_prompt(job)
    assert payload["source_catalog"]
    assert "master-profile facts" in truth_rules({"stage": "edit_cv_sections"})
    assert truth_rules({"stage": "tailor_resume"}).startswith("MASTER PROFILE")


def test_custom_instructions_control_existing_employment_fact_conditions(profile):
    from app.services.cv_content_policy import cv_catalog
    fact = profile.experience[0].facts[0]
    profile.experience[0].cv_fact_conditions[fact] = ['migrations']
    job = custom_job()
    assert job.job_description == ''
    assert not any(e.kind == 'experience' and e.text == fact for e in cv_catalog(profile, job).prompt_entries)
    job.custom_instructions += ' Emphasize migrations.'
    assert any(e.kind == 'experience' and e.text == fact for e in cv_catalog(profile, job).prompt_entries)


@pytest.mark.parametrize("provider_type", [OllamaProvider, GeminiProvider])
async def test_custom_generation_skips_job_analysis(profile, monkeypatch, provider_type):
    provider = provider_type(Settings(_env_file=None))
    if provider.name == "gemini":
        provider.api_key = "synthetic"
    monkeypatch.setattr(provider, "analyze_job", AsyncMock(side_effect=AssertionError("No offer analysis")))
    calls = []
    async def chat(model, prompt, payload, **kwargs):
        calls.append((prompt, payload))
        return TailoredResume(headline="Support", professional_summary="", core_skills=[], experience=[])
    monkeypatch.setattr(provider, "_chat", chat)
    router = ResilientAIProvider({provider.name: provider}, default_primary=provider.name, default_fallback="none")
    response = await router.tailor(custom_job(), profile, requested_provider=provider.name)
    assert len(calls) == 1
    assert calls[0][1]["custom_instructions"].startswith("Emphasize")
    assert response.analysis.is_reliable and not response.analysis.requirements
    provider.analyze_job.assert_not_awaited()


def test_custom_master_generation_removes_invented_qualifications(profile, resume_payload):
    job = custom_job()
    resume = TailoredResume.model_validate(resume_payload)
    result = FactValidator(profile).validate(resume, job)
    assert "Kubernetes" not in str(result.resume.model_dump())
    assert result.resume.experience[0].dates == "2022 – Present"
    assert result.resume.experience[0].title == profile.experience[0].title


def test_custom_section_preferences_reach_existing_pdf_and_docx_templates(profile, resume_payload, tmp_path):
    from pathlib import Path

    from bs4 import BeautifulSoup
    from docx import Document

    from app.services.docx_generator import _generate_docx_local
    from app.services.pdf_generator import PDFGenerator
    job = custom_job()
    resume = FactValidator(profile).validate(TailoredResume.model_validate(resume_payload), job).resume
    resume.section_order = ['experience', 'summary']
    resume.excluded_sections = ['skills']
    html = PDFGenerator(Path('templates'), Path('static'), tmp_path, isolated=False).render_html(resume, profile, template_name='ats_classic')
    headings = [h.get_text() for h in BeautifulSoup(html, 'html.parser').select('h2')]
    assert headings.index('Experience') < headings.index('Professional Summary')
    assert 'Core Skills' not in headings
    output = tmp_path / 'custom.docx'
    _generate_docx_local(resume, profile, output, template_name='ats_classic')
    headings = [p.text for p in Document(output).paragraphs if p.style.name == 'Heading 1']
    assert headings.index('Experience') < headings.index('Professional Summary')
    assert 'Core Skills' not in headings


def test_long_cv_distinct_metrics_skip_inapplicable_deduplication(monkeypatch):
    from app.services import content_quality
    compare = content_quality._subsumed
    calls = []
    def counted(first, second):
        calls.append((first, second))
        return compare(first, second)
    monkeypatch.setattr(content_quality, '_subsumed', counted)
    original = ' '.join(f'Resolved documented incident {i}.' for i in range(600))
    assert content_quality.deduplicate_text(original) == original
    assert calls == []


@pytest.mark.parametrize("text", [SOURCE, SOURCE.replace("\n", "\r\n"), "Name\nCustom heading\nText with no known sections"])
def test_import_sections_are_lossless(text):
    assert "".join(s["text"] for s in split_sections(text)) == text


def test_uppercase_employer_is_not_misclassified_as_a_section():
    sections = split_sections('ADA EXAMPLE\nProfessional Summary\nSupport specialist.\nExperience\nMOTOROLA SOLUTIONS\nEngineer | 2022–Present\nEducation\nBSc | 2021\n')
    assert [s['title'] for s in sections] == ['Header / unclassified content', 'Professional Summary', 'Experience', 'Education']
    assert 'MOTOROLA SOLUTIONS' in sections[2]['text']


@pytest.mark.parametrize("text", ["", "  ", "a" * 60001, "Name\x00CV"])
def test_invalid_import_is_rejected(tmp_path, text):
    with pytest.raises(ValueError):
        CVEditorStore(tmp_path).create("Example", text)


def summary_edit(text="Professional Summary\nTroubleshoots SQL incidents.\n\n"):
    return EditProposal(replacements=[SectionReplacement(section_id="s1", text=text)])


def test_summary_edit_preserves_unrelated_sections_exactly():
    sections = split_sections(SOURCE)
    edited = apply_proposal(sections, summary_edit(), ["s1"])
    assert edited[edited.index("Experience"):] == SOURCE[SOURCE.index("Experience"):]
    assert edited[:edited.index("Professional Summary")] == SOURCE[:SOURCE.index("Professional Summary")]


def test_out_of_scope_and_invented_metrics_are_rejected():
    with pytest.raises(ProviderResponseError, match="outside"):
        apply_proposal(split_sections(SOURCE), summary_edit(), ["s2"])
    with pytest.raises(ProviderResponseError, match="dates or numbers"):
        apply_proposal(split_sections(SOURCE), summary_edit("Professional Summary\nAchieved 99% improvement.\n"), ["s1"])


async def test_independent_review_fails_closed_on_unsupported_qualification():
    chat = AsyncMock(side_effect=[summary_edit("Professional Summary\nAWS certified engineer.\n"),
        EditReview(supported=False, scope_preserved=True, issues=["Do you hold an AWS certification? Please provide the confirmed qualification."])])
    proposal = await edit_with_chat(chat, {"stage": "edit_cv_sections", "sections": split_sections(SOURCE),
        "selected_sections": ["s1"], "instructions": "Highlight cloud qualifications", "confirmed_additions": ""})
    assert not proposal.replacements and "AWS" in proposal.clarification
    assert chat.await_count == 2


async def test_clarification_does_not_publish_changes():
    chat = AsyncMock(return_value=EditProposal(clarification="Please confirm your experience.", replacements=summary_edit().replacements))
    proposal = await edit_with_chat(chat, {"sections": split_sections(SOURCE), "selected_sections": ["s1"]})
    assert not proposal.replacements


@pytest.mark.parametrize('provider_type', [OllamaProvider, GeminiProvider])
async def test_section_editor_uses_existing_http_provider_with_separate_truth_source(provider_type):
    import json
    calls = []
    def respond(request):
        body = json.loads(request.content)
        if provider_type is OllamaProvider:
            payload = json.loads(body['messages'][1]['content'])
            system = body['messages'][0]['content']
        else:
            payload = json.loads(body['contents'][0]['parts'][0]['text'])
            system = body['systemInstruction']['parts'][0]['text']
        calls.append(payload)
        assert 'MASTER PROFILE IS THE ONLY' not in system
        assert 'Jan Test' not in json.dumps(payload)
        result = summary_edit().model_dump(mode='json') if payload['stage'] == 'edit_cv_sections' else {
            'supported': True, 'scope_preserved': True, 'issues': []}
        if provider_type is OllamaProvider:
            return httpx.Response(200, json={'message': {'content': json.dumps(result)}, 'done': True})
        return httpx.Response(200, json={'candidates': [{'content': {'parts': [{'text': json.dumps(result)}]}}]})
    provider = provider_type(Settings(_env_file=None, gemini_api_key='synthetic'), transport=httpx.MockTransport(respond))
    proposal = await provider.edit_cv({'stage': 'edit_cv_sections', 'sections': split_sections(SOURCE),
        'selected_sections': ['s1'], 'instructions': 'Improve only Professional Summary.', 'confirmed_additions': ''})
    assert len(calls) == 2
    assert 'Troubleshoots SQL incidents.' in apply_proposal(split_sections(SOURCE), proposal, ['s1'])


async def test_section_editor_fallback_preserves_payload():
    first = SimpleNamespace(edit_cv=AsyncMock(side_effect=ProviderUnavailableError('synthetic unavailable')))
    second = SimpleNamespace(edit_cv=AsyncMock(return_value=summary_edit()))
    provider = ResilientAIProvider({'ollama': first, 'gemini': second})
    payload = {'stage': 'edit_cv_sections', 'sections': split_sections(SOURCE),
               'selected_sections': ['s1'], 'instructions': 'Improve summary', 'confirmed_additions': ''}
    assert await provider.edit_cv(payload) == summary_edit()
    first.edit_cv.assert_awaited_once_with(payload)
    second.edit_cv.assert_awaited_once_with(payload)


async def test_custom_task_keeps_lazy_model_selection_and_generation_only_fallback(profile, resume_payload, tmp_path, monkeypatch):
    from app.services.storage import Storage

    def unavailable(request):
        raise httpx.ConnectError('Synthetic Ollama discovery failure', request=request)
    primary = OllamaProvider(Settings(_env_file=None), transport=httpx.MockTransport(unavailable))
    fallback = GeminiProvider(Settings(_env_file=None, gemini_api_key='synthetic'))
    monkeypatch.setattr(fallback, 'tailor_resume', AsyncMock(return_value=TailoredResume.model_validate(resume_payload)))
    for target in (primary, fallback):
        monkeypatch.setattr(target, 'analyze_job', AsyncMock(side_effect=AssertionError('Custom CV has no job analysis')))
    provider = ResilientAIProvider({'ollama': primary, 'gemini': fallback})
    state = SimpleNamespace(profile=profile, skills_bank=None, provider=provider,
        settings=SimpleNamespace(csrf_secret='synthetic-review-key'), storage=Storage(tmp_path))
    manager = TaskManager(tmp_path)
    task = manager.create_task(TaskKind.TAILOR, {'job': custom_job().model_dump(mode='json'),
        'llm_provider': 'ollama', 'ollama_model': 'synthetic-selected', 'ollama_context': 65536}, 'custom-lazy-selection')
    await manager._execute_tailor(task, state)
    completed = manager.get_task(task.task_id)
    assert completed.status == TaskStatus.COMPLETED
    assert completed.result['fallback_used'] and completed.result['provider_used'] == 'gemini'
    job, response = state.storage.load_draft(completed.result['draft_id'])
    assert job.workflow == 'custom' and job.job_description == ''
    assert not response.analysis.requirements
    fallback.tailor_resume.assert_awaited_once()
    assert primary.model != 'synthetic-selected'  # shared provider was not mutated


async def test_custom_generation_final_review_uses_instructions_not_a_job_description(profile):
    import json

    from app.services.ollama_draft_review import review_draft, valid_bullet_attestation
    from tests.test_ollama_draft_review import inputs, review_response
    job = custom_job()
    _, _, resume = inputs(profile)
    calls = []
    def respond(request):
        body = json.loads(request.content)
        calls.append(body)
        return review_response(body)
    provider = OllamaProvider(Settings(_env_file=None, csrf_secret='synthetic-review-key'),
                              transport=httpx.MockTransport(respond))
    checked = await review_draft(provider, resume, custom_analysis(job), job, profile)
    assert len(calls) == 2
    for body in calls:
        payload = json.loads(body['messages'][1]['content'])
        assert payload['workflow'] == 'custom'
        assert payload['custom_instructions'] == job.custom_instructions
        assert payload['job_description'] == ''
    assert 'custom_instructions' in calls[-1]['messages'][0]['content']
    assert valid_bullet_attestation(checked.experience[0].bullets[0], checked, profile, 'synthetic-review-key')


def test_version_accept_discard_undo_restore_and_duplicate(tmp_path):
    store = CVEditorStore(tmp_path)
    identifier = store.create("Support CV", SOURCE)
    first = apply_proposal(split_sections(SOURCE), summary_edit(), ["s1"])
    store.propose(identifier, 0, first, "Improve summary")
    assert store.load(identifier)["current"] == SOURCE
    store.update(identifier, 0, "accept")
    assert store.load(identifier)["current"] == first
    second = first.replace("Resolved customer incidents", "Resolved SQL customer incidents")
    store.propose(identifier, 1, second, "Improve experience")
    store.update(identifier, 1, "discard")
    assert store.load(identifier)["current"] == first
    store.propose(identifier, 2, second, "Improve experience")
    store.update(identifier, 2, "accept")
    assert store.load(identifier)["current"] == second
    store.update(identifier, 3, "undo")
    assert store.load(identifier)["current"] == first
    duplicate = store.duplicate(identifier)
    assert store.load(duplicate)["current"] == first
    assert store.load(duplicate)["original"] == SOURCE
    store.update(identifier, 4, "restore")
    assert store.load(identifier)["current"] == SOURCE
    assert store.load(duplicate)["current"] == first
    store.delete(duplicate)
    assert len(store.list()) == 1


def test_stale_proposal_cannot_replace_newer_manual_edits(tmp_path):
    store = CVEditorStore(tmp_path)
    identifier = store.create("CV", SOURCE)
    store.update(identifier, 0, "manual", text=SOURCE + "Additional details\n")
    with pytest.raises(ValueError, match="another tab"):
        store.propose(identifier, 0, SOURCE, "Stale prompt")
    assert store.load(identifier)["current"].endswith("Additional details\n")
    assert store.load(identifier)["original"] == SOURCE


def test_renaming_does_not_discard_a_pending_proposal(tmp_path):
    store = CVEditorStore(tmp_path)
    identifier = store.create('CV', SOURCE)
    revised = apply_proposal(split_sections(SOURCE), summary_edit(), ['s1'])
    store.propose(identifier, 0, revised, 'Improve summary')
    store.update(identifier, 0, 'rename', name='Better name')
    assert store.load(identifier)['pending']['text'] == revised
    store.update(identifier, 1, 'accept')
    assert store.load(identifier)['current'] == revised


async def test_consecutive_ai_tasks_use_latest_accepted_version(tmp_path):
    store = CVEditorStore(tmp_path)
    identifier = store.create("CV", SOURCE)
    provider = SimpleNamespace(edit_cv=AsyncMock(return_value=summary_edit()))
    state = SimpleNamespace(settings=SimpleNamespace(data_dir=tmp_path), provider=provider)
    manager = TaskManager(tmp_path)
    for revision in (0, 1):
        task = manager.create_task(TaskKind.CUSTOM_EDIT, {"cv_id": identifier, "revision": revision,
            "selected_sections": ["s1"], "instructions": "Improve the summary", "llm_provider": "ollama"}, f"key{revision}")
        await manager._execute_custom_edit(task, state)
        assert manager.get_task(task.task_id).status == TaskStatus.COMPLETED
        store.update(identifier, revision, "accept")
    second_payload = provider.edit_cv.await_args_list[1].args[0]
    assert "Troubleshoots SQL incidents." in str(second_payload["sections"])
    assert "master" not in str(second_payload)
    assert store.load(identifier)["original"] == SOURCE


async def test_ai_failure_keeps_saved_document_unchanged(tmp_path):
    store = CVEditorStore(tmp_path)
    identifier = store.create("CV", SOURCE)
    before = store.path(identifier).read_bytes()
    state = SimpleNamespace(settings=SimpleNamespace(data_dir=tmp_path),
                            provider=SimpleNamespace(edit_cv=AsyncMock(side_effect=ProviderUnavailableError("Unavailable"))))
    manager = TaskManager(tmp_path)
    task = manager.create_task(TaskKind.CUSTOM_EDIT, {"cv_id": identifier, "revision": 0,
        "selected_sections": ["s1"], "instructions": "Improve summary"}, "failure")
    with pytest.raises(ProviderUnavailableError):
        await manager._execute_custom_edit(task, state)
    assert store.path(identifier).read_bytes() == before


def test_import_routes_csrf_manual_edit_library_and_master_independence():
    with TestClient(app) as client:
        sign_in(client)
        master = app.state.settings.master_profile_path
        before = master.read_bytes()
        assert client.get("/").status_code == 200
        assert client.post("/custom-cvs/import", data={"csrf_token": "wrong", "pasted_cv": SOURCE}).status_code == 403
        token = client.cookies.get("csrf_token")
        response = client.post("/custom-cvs/import", data={"csrf_token": token, "cv_name": "My support CV", "pasted_cv": SOURCE}, follow_redirects=False)
        assert response.status_code == 303, response.text
        url = response.headers["location"]
        assert "My support CV" in client.get(url).text
        assert "My support CV" in client.get("/custom-cvs/").text
        assert client.post(url + "/apply", data={"csrf_token": token, "revision": 0, "instructions": "Improve summary"}).status_code == 422
        response = client.post(url + "/manual", data={"csrf_token": token, "revision": 0, "text": SOURCE.replace("I troubleshoot", "I resolve")}, follow_redirects=False)
        assert response.status_code == 303
        reopened = client.get(url)
        assert "I resolve" in reopened.text and "I troubleshoot" in reopened.text  # recoverable original
        assert client.post(url + "/manual", data={"csrf_token": token, "revision": 0, "text": "stale"}).status_code == 409
        assert master.read_bytes() == before
        assert client.get("/custom-cvs/../../master_profile.json").status_code == 404


def test_file_import_reuses_isolated_extraction_and_pasted_text_takes_precedence(monkeypatch):
    calls = []
    def extract(operation, payload, **kwargs):
        calls.append((operation, payload))
        return SOURCE
    monkeypatch.setattr('app.routes.custom_cvs.process', extract)
    with TestClient(app) as client:
        sign_in(client)
        token = client.cookies.get('csrf_token')
        response = client.post('/custom-cvs/import', data={'csrf_token': token, 'cv_name': 'File CV'},
            files={'cv_file': ('cv.docx', b'synthetic extraction fixture')}, follow_redirects=False)
        assert response.status_code == 303
        assert calls[0][0] == 'reference' and calls[0][1]['suffix'] == '.docx'
        assert SOURCE in CVEditorStore(app.state.settings.data_dir).load(response.headers['location'].split('/')[-1])['original']
        response = client.post('/custom-cvs/import', headers={'Accept': 'application/json'},
            data={'csrf_token': token, 'pasted_cv': SOURCE}, files={'cv_file': ('ignored.pdf', b'ignored')})
        assert response.status_code == 202 and response.json()['view_url'].startswith('/custom-cvs/')
        assert len(calls) == 1
        assert client.post('/custom-cvs/import', data={'csrf_token': token},
            files={'cv_file': ('bad.exe', b'invalid')}).status_code == 422


def test_ai_editor_http_queue_review_accept_iterate_and_undo(monkeypatch):
    with TestClient(app) as client:
        sign_in(client)
        store = CVEditorStore(app.state.settings.data_dir)
        identifier = store.create('Iterative CV', SOURCE)
        base = f'/custom-cvs/{identifier}'
        token = client.cookies.get('csrf_token')
        master_before = app.state.settings.master_profile_path.read_bytes()
        ai = AsyncMock(return_value=summary_edit())
        monkeypatch.setattr(app.state.provider, 'edit_cv', ai)
        first = ''
        for revision in (0, 1):
            if revision:
                ai.return_value = summary_edit('Professional Summary\nSQL incident troubleshooting.\n\n')
            response = client.post(base + '/apply', data={'csrf_token': token, 'revision': revision,
                'instructions': 'Improve only the summary.', 'sections': ['s1']}, follow_redirects=False)
            assert response.status_code == 303
            client.portal.call(app.state.task_manager._queue.join)
            task_id = response.headers['location'].split('/')[-2]
            assert app.state.task_manager.get_task(task_id).status == TaskStatus.COMPLETED
            assert store.load(identifier)['current'] == (SOURCE if revision == 0 else first)
            assert '<ins>' in client.get(base).text
            assert client.post(base + '/accept', data={'csrf_token': token, 'revision': revision}, follow_redirects=False).status_code == 303
            if revision == 0:
                first = store.load(identifier)['current']
        assert ai.await_count == 2
        assert 'Troubleshoots SQL incidents.' in str(ai.await_args_list[1].args[0]['sections'])
        assert store.load(identifier)['current'].endswith(SOURCE[SOURCE.index('Experience'):])
        assert client.post(base + '/undo', data={'csrf_token': token, 'revision': 2}, follow_redirects=False).status_code == 303
        assert store.load(identifier)['current'] == first
        assert store.load(identifier)['original'] == SOURCE
        assert app.state.settings.master_profile_path.read_bytes() == master_before


def test_analyze_endpoint_custom_without_company_role_or_job():
    with TestClient(app) as client:
        sign_in(client)
        response = client.post("/analyze", headers={"Accept": "application/json"}, data={
            "csrf_token": client.cookies.get("csrf_token"), "workflow": "custom", "cv_name": "Technical Support CV",
            "job_description": "Emphasize verified SQL and troubleshooting experience.", "llm_provider": "mock"})
        assert response.status_code == 202, response.text
        task = app.state.task_manager.get_task(response.json()["task_id"])
        assert task.payload["job"]["job_description"] == ""
        assert task.payload["job"]["custom_instructions"].startswith("Emphasize")
        assert task.payload["job"]["workflow"] == "custom"


def test_custom_draft_accept_save_reopen_export_and_regenerate(resume_payload):
    from bs4 import BeautifulSoup

    from app.models.resume import WorkflowResponse
    with TestClient(app) as client:
        sign_in(client)
        before = app.state.settings.master_profile_path.read_bytes()
        job = custom_job(cv_name='Support version')
        checked = FactValidator(app.state.profile, app.state.skills_bank, custom_analysis(job),
                                review_key=app.state.settings.csrf_secret).validate(TailoredResume.model_validate(resume_payload), job).resume
        draft = app.state.storage.save_draft(job, WorkflowResponse(analysis=custom_analysis(job), resume=checked))
        preview = client.get(f'/analysis/{draft}')
        assert preview.status_code == 200, preview.text
        assert 'Accept and save CV' in preview.text
        soup = BeautifulSoup(preview.text, 'html.parser')
        assert soup.select_one('form[data-document-export]') is not None
        busy = client.get(f'/analysis/{draft}?export_busy=1')
        assert busy.status_code == 200 and 'Your draft is saved' in busy.text
        token = soup.select_one('[name="summary_token"]')
        data = {'csrf_token': client.cookies.get('csrf_token'), 'cv_name': 'Accepted custom CV',
                'template_name': 'ats_classic', 'summary_confirmed': 'on'}
        if token:
            data['summary_token'] = token['value']
        accepted = client.post(f'/generate/{draft}', data=data, follow_redirects=False)
        assert accepted.status_code == 303, accepted.text
        slug = accepted.headers['location'].split('/')[-1]
        assert 'Accepted custom CV' in client.get(f'/preview/{slug}').text
        assert client.get(f'/download/{slug}/pdf').content.startswith(b'%PDF-')
        assert client.get(f'/download/{slug}/docx').content.startswith(b'PK')
        revised = client.get(f'/?from_cv={slug}')
        assert revised.status_code == 200 and job.custom_instructions in revised.text
        assert 'Accepted custom CV' in client.get('/history').text
        duplicated = client.post(f'/preview/{slug}/version/duplicate', data={'csrf_token': client.cookies.get('csrf_token')}, follow_redirects=False)
        assert duplicated.status_code == 303
        assert duplicated.headers['location'] != accepted.headers['location']
        editor_copy = client.post(f'/preview/{slug}/ai-editor', data={'csrf_token': client.cookies.get('csrf_token')}, follow_redirects=False)
        assert editor_copy.status_code == 303
        assert 'Current CV' in client.get(editor_copy.headers['location']).text
        assert app.state.settings.master_profile_path.read_bytes() == before


@pytest.mark.parametrize('kind', ['pdf', 'docx'])
def test_text_cv_routes_export_accepted_content_only(kind):
    from io import BytesIO

    from docx import Document
    from pypdf import PdfReader
    with TestClient(app) as client:
        sign_in(client)
        store = CVEditorStore(app.state.settings.data_dir)
        identifier = store.create('Original source', SOURCE)
        proposal = SOURCE.replace('I troubleshoot SQL incidents.', 'UNACCEPTED PROPOSAL')
        store.propose(identifier, 0, proposal, 'Example proposal')
        page = client.get(f'/custom-cvs/{identifier}')
        assert page.status_code == 200 and '<ins>' in page.text and '<del>' in page.text
        response = client.get(f'/custom-cvs/{identifier}/export/{kind}')
        assert response.status_code == 200, response.text[:500]
        if kind == 'pdf':
            content = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages)
        else:
            content = '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        assert 'UNACCEPTED PROPOSAL' not in content
        assert 'I troubleshoot SQL incidents.' in content
        assert '2022' in content and '2021' in content
        assert content.count('Experience') == 1
        assert 'Jan Test' not in content  # master identity must not leak into pasted-CV exports
