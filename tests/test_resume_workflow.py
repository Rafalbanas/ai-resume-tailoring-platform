import json
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.core.config import Settings
from app.models.candidate import CandidateProfile
from app.models.job import JobAnalysis, JobRequest
from app.models.resume import TailoredResume, WorkflowResponse
from app.models.skills import SkillEvidence, SkillsDocument, VerifiedSkill
from app.services.analysis_validator import AnalysisValidator
from app.services.docx_generator import generate_docx
from app.services.fact_catalog import FactCatalog
from app.services.fact_validator import FactValidator
from app.services.pdf_generator import PDFGenerator
from app.services.profile_loader import ProfileConfigurationError, load_master_profile
from app.services.resume_editor import apply_resume_edits
from app.services.skills_bank import SkillsBank
from app.services.storage import Storage


@pytest.fixture
def verified_profile() -> CandidateProfile:
    return CandidateProfile.model_validate(
        {
            "personal": {
                "name": "Rafał Test",
                "location": "Katowice, Poland",
                "email": "rafal@example.test",
            },
            "summary_facts": [
                "IT professional with more than three years of support and operations experience.",
                "Uses English at B2 level in daily professional work.",
            ],
            "skills": {
                "systems": ["Windows Server", "Active Directory", "Linux"],
                "automation": ["PowerShell", "Python", "Bash (basic)"],
                "languages": ["English B2"],
            },
            "experience": [
                {
                    "company": "Motorola Solutions",
                    "title": "Technical Support Engineer (Tier 2)",
                    "start": "Apr 2026",
                    "end": "Present",
                    "facts": [
                        "Troubleshoots Windows and network incidents for partners across EMEA.",
                        "Automates selected operational tasks using PowerShell and Python.",
                    ],
                }
            ],
            "education": [
                {
                    "institution": "Silesian University of Technology",
                    "qualification": "Bachelor's Degree in Automation and Robotics",
                    "dates": "2022",
                },
                {
                    "institution": "University of Economics in Katowice",
                    "qualification": "MSc in Computer Science",
                    "dates": "2026",
                },
            ],
            "projects": [
                {
                    "name": "Linux VPS",
                    "description": "Deployment and maintenance of self-hosted services on a Linux VPS.",
                    "technologies": ["Linux"],
                    "facts": ["Maintains self-hosted services using SSH and Linux administration."],
                }
            ],
            "certifications": [],
        }
    )


@pytest.fixture
def platform_job() -> JobRequest:
    return JobRequest(
        company="Example Finance",
        role="Platform Engineer",
        job_description=(
            "Platform Engineer requiring PostgreSQL, Azure Data Factory, GitLab CI/CD, Elasticsearch, "
            "a Bachelor's degree, good English, Linux, Python and operational automation."
        ),
    )


def job_analysis() -> JobAnalysis:
    return JobAnalysis(
        company="Example Finance",
        role="Platform Engineer",
        strong_matches=["Python"],
        partial_matches=["Linux"],
        missing_requirements=["PostgreSQL"],
        supported_keywords=["Python"],
        unsupported_keywords=["PostgreSQL"],
        recommendation="REASONABLE_STRETCH",
        match_level="MEDIUM",
        reasoning_summary="Validated match.",
    )


def generated_resume(profile: CandidateProfile, job: JobRequest) -> TailoredResume:
    empty = TailoredResume(
        headline=job.role,
        professional_summary="",
        summary_source_fact_ids=[],
        core_skills=[],
        experience=[],
        education=[],
        projects=[],
        certifications=[],
    )
    return FactValidator(profile).validate(empty, job).resume


def test_production_profile_never_falls_back_to_example(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path, profile_mode="production")
    with pytest.raises(ProfileConfigurationError, match="never a fallback"):
        load_master_profile(settings)

    example = json.loads(Path("data/master_profile.example.json").read_text(encoding="utf-8"))
    (tmp_path / "master_profile.json").write_text(json.dumps(example), encoding="utf-8")
    with pytest.raises(ProfileConfigurationError, match="example or template"):
        load_master_profile(settings)


def test_platform_analysis_rejects_example_technologies_and_recovers_real_requirements(verified_profile, platform_job):
    python_source = FactCatalog(verified_profile).direct_sources("Python")
    analysis = JobAnalysis(
        company=platform_job.company,
        role=platform_job.role,
        strong_matches=["PostgreSQL", "Azure Data Factory", "GitLab CI/CD", "Elasticsearch"],
        partial_matches=[],
        missing_requirements=["Bachelor's degree", "Good English"],
        match_sources={
            "PostgreSQL": python_source,
            "Azure Data Factory": python_source,
            "GitLab CI/CD": python_source,
            "Elasticsearch": python_source,
        },
        supported_keywords=["PostgreSQL", "Azure Data Factory", "GitLab CI/CD", "Elasticsearch"],
        unsupported_keywords=[],
        recommendation="APPLY",
        match_level="HIGH",
        reasoning_summary="The master profile is a template.",
    )

    result = AnalysisValidator(verified_profile).validate(analysis)

    for unsupported in ("PostgreSQL", "Azure Data Factory", "GitLab CI/CD", "Elasticsearch"):
        assert unsupported not in result.strong_matches
        assert unsupported in result.missing_requirements
    assert "Bachelor's degree" not in result.missing_requirements
    assert "Good English" not in result.missing_requirements
    assert "Bachelor's degree" in result.strong_matches
    assert "Good English" in result.strong_matches
    assert "template" not in result.reasoning_summary.casefold()
    assert all(result.match_sources[item] for item in result.strong_matches)


def test_truth_lock_preserves_supported_source_referenced_paraphrase(verified_profile, platform_job):
    catalog = FactCatalog(verified_profile)
    summary_id = catalog.direct_sources("support operations")[0]
    experience_id = next(
        entry.source_id
        for entry in catalog.prompt_entries
        if entry.kind == "experience" and "Automates selected" in entry.text
    )
    candidate = TailoredResume(
        headline="Platform Engineer",
        professional_summary="IT professional with over three years in support and operations.",
        summary_source_fact_ids=[summary_id],
        core_skills=["Python", "Bash"],
        experience=[
            {
                "company": "Different model wording",
                "title": "Changed title",
                "dates": "",
                "bullets": [
                    {
                        "text": "Automates operational tasks with Python and PowerShell.",
                        "source_fact_ids": [experience_id],
                    }
                ],
            }
        ],
        education=[],
        projects=[],
        certifications=[],
    )

    result = FactValidator(verified_profile).validate(candidate, platform_job)

    assert result.resume.professional_summary == candidate.professional_summary
    assert result.resume.experience[0].bullets[0].text == candidate.experience[0].bullets[0].text
    assert result.resume.experience[0].company == "Motorola Solutions"
    assert result.resume.core_skills == ["Python", "Bash (basic)"]


def test_project_titles_are_verified_separately_from_technologies(verified_profile, platform_job, tmp_path):
    payload = verified_profile.model_dump(mode="json")
    payload["projects"] = [
        {
            "name": "Linux VPS / banas.dev",
            "description": "Deployment and maintenance of verified self-hosted services on a Linux VPS.",
            "technologies": ["Linux"],
            "facts": ["Maintains verified services through SSH and Linux administration."],
        },
        {
            "name": "Python ML thesis project",
            "description": "MSc thesis project comparing multiple machine-learning models in Python.",
            "technologies": ["Python"],
            "facts": ["Built a verified Python data preprocessing and model comparison pipeline."],
        },
    ]
    profile = CandidateProfile.model_validate(payload)
    skills = SkillsDocument(
        skills=[
            VerifiedSkill(
                id="skill_linux",
                name="Linux",
                category="Linux",
                level="hands_on",
                verified=True,
                allowed_in_cv=True,
                cv_wording="Hands-on Linux administration",
                evidence=[
                    SkillEvidence(
                        source_type="manual_verified",
                        source_id="manual:linux",
                        description="Verified Linux test evidence",
                    )
                ],
            ),
            VerifiedSkill(
                id="skill_python",
                name="Python",
                category="Automation / Scripting",
                level="intermediate",
                verified=True,
                allowed_in_cv=True,
                cv_wording="Python automation",
                evidence=[
                    SkillEvidence(
                        source_type="manual_verified",
                        source_id="manual:python",
                        description="Verified Python test evidence",
                    )
                ],
            ),
        ]
    )
    skills_path = tmp_path / "skills.json"
    skills_path.write_text(skills.model_dump_json(), encoding="utf-8")
    bank = SkillsBank(skills_path, profile)
    catalog = FactCatalog(profile)

    projects = []
    for index, source in enumerate(profile.projects):
        source_ids = [
            entry.source_id
            for entry in catalog.prompt_entries
            if entry.kind == "project" and entry.owner_index == index
        ]
        projects.append(
            {
                "name": source.name,
                "description": source.description,
                "technologies": [*source.technologies, "PostgreSQL"] if index == 1 else source.technologies,
                "source_fact_ids": source_ids,
            }
        )
    candidate_payload = generated_resume(profile, platform_job).model_dump(mode="json")
    candidate_payload["projects"] = projects
    candidate = TailoredResume.model_validate(candidate_payload)
    result = FactValidator(profile, bank).validate(candidate, platform_job)

    by_name = {project.name: project for project in result.resume.projects}
    assert "Linux VPS / banas.dev" in by_name
    assert "Python ML thesis project" in by_name
    assert by_name["Linux VPS / banas.dev"].technologies == ["Hands-on Linux administration"]
    assert by_name["Python ML thesis project"].technologies == ["Python automation"]
    assert all(project.name not in project.technologies for project in result.resume.projects)
    assert any("PostgreSQL" in warning for warning in result.warnings)
    assert not any("Linux VPS / banas.dev:" in warning for warning in result.warnings)


def test_generated_resume_is_complete_and_has_no_candidate_placeholder(verified_profile, platform_job):
    validation = FactValidator(verified_profile).validate(
        TailoredResume(
            headline=platform_job.role,
            professional_summary="",
            summary_source_fact_ids=[],
            core_skills=[],
            experience=[],
            education=[],
            projects=[],
            certifications=[],
        ),
        platform_job,
    )
    resume = validation.resume
    html = PDFGenerator(Path("templates"), Path("static")).render_html(resume, verified_profile)

    assert not validation.warnings
    assert resume.professional_summary
    assert resume.core_skills
    assert resume.experience
    assert resume.education
    assert resume.projects
    assert "Rafał Test" in html
    assert "Candidate Name" not in html


def test_edit_save_reject_unsupported_and_reset(verified_profile, platform_job, tmp_path):
    storage = Storage(tmp_path)
    original = generated_resume(verified_profile, platform_job)
    slug, _ = storage.save_application(platform_job, job_analysis(), original, [])

    safe_form = {
        "headline": "Windows Platform Engineer",
        "professional_summary": "IT professional with over three years in support and operations.",
        "core_skills": "Python\nLinux",
        "experience_0_bullet_0": "Troubleshoots Windows and network incidents across EMEA.",
        "experience_0_bullet_1": original.experience[0].bullets[1].text,
        "project_0_description": original.projects[0].description,
        "project_0_technologies": "Linux",
    }
    edited = apply_resume_edits(original, safe_form)
    safe_result = FactValidator(verified_profile).validate(edited, platform_job)
    storage.save_current_resume(slug, safe_result.resume, safe_result.warnings)
    _, saved, generated, _ = storage.load_application(slug)
    assert saved.headline == "Platform Engineer | Python | Linux"
    assert saved.experience[0].bullets[0].text == safe_form["experience_0_bullet_0"]

    unsupported_form = dict(safe_form)
    unsupported_form["core_skills"] = "Python\nKubernetes"
    unsupported_form["experience_0_bullet_0"] = "Administered Kubernetes clusters in production."
    rejected = FactValidator(verified_profile).validate(apply_resume_edits(saved, unsupported_form), platform_job)
    assert "Kubernetes" not in rejected.resume.model_dump_json()
    assert rejected.resume.experience[0].bullets[0].text != unsupported_form["experience_0_bullet_0"]
    assert rejected.warnings

    storage.save_current_resume(slug, generated, [])
    _, reset, _, _ = storage.load_application(slug)
    assert reset == original


def test_delete_single_and_clear_history_preserve_profile_and_references(verified_profile, platform_job, tmp_path):
    storage = Storage(tmp_path)
    resume = generated_resume(verified_profile, platform_job)
    workflow = WorkflowResponse(analysis=job_analysis(), resume=resume)
    master = tmp_path / "master_profile.json"
    reference = tmp_path / "reference_cvs" / "reference.pdf"
    master.write_text("private profile", encoding="utf-8")
    reference.parent.mkdir()
    reference.write_bytes(b"%PDF-reference")

    draft_id = storage.save_draft(platform_job, workflow)
    slug, folder = storage.save_application(platform_job, workflow.analysis, resume, [], draft_id=draft_id)
    (folder / "resume.pdf").write_bytes(b"%PDF")
    (folder / "resume.docx").write_bytes(b"DOCX")
    storage.delete_application(slug)
    assert not folder.exists()
    assert not (storage.drafts_dir / f"{draft_id}.json").exists()
    assert master.exists()
    assert reference.exists()

    first, _ = storage.save_application(platform_job, workflow.analysis, resume, [])
    second, _ = storage.save_application(platform_job, workflow.analysis, resume, [])
    assert {item["slug"] for item in storage.history()} == {first, second}
    assert storage.clear_history() == 2
    assert not storage.history()
    assert master.exists()
    assert reference.exists()


def test_pdf_and_docx_use_current_edited_resume(verified_profile, platform_job, tmp_path):
    storage = Storage(tmp_path)
    original = generated_resume(verified_profile, platform_job)
    slug, _ = storage.save_application(platform_job, job_analysis(), original, [])
    form = {
        "headline": "Windows Platform Engineer",
        "professional_summary": original.professional_summary,
        "core_skills": "\n".join(original.core_skills),
        "experience_0_bullet_0": original.experience[0].bullets[0].text,
        "experience_0_bullet_1": original.experience[0].bullets[1].text,
        "project_0_description": original.projects[0].description,
        "project_0_technologies": "Linux",
    }
    edited = FactValidator(verified_profile).validate(apply_resume_edits(original, form), platform_job).resume
    folder = storage.save_current_resume(slug, edited, [])
    PDFGenerator(Path("templates"), Path("static"), tmp_path).generate(edited, verified_profile, folder / "resume.pdf")
    generate_docx(edited, verified_profile, folder / "resume.docx")

    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(folder / "resume.pdf").pages)
    document = Document(folder / "resume.docx")
    docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Platform Engineer" in pdf_text
    assert "Python" in pdf_text
    assert "Platform Engineer | Linux | Python" in docx_text
    assert storage.artifact(slug, "resume.pdf").is_file()
    assert storage.artifact(slug, "resume.docx").is_file()


def test_two_most_relevant_verified_projects_are_selected(verified_profile, platform_job):
    payload = verified_profile.model_dump(mode="json")
    payload["projects"].extend(
        [
            {
                "name": "Unrelated Hobby",
                "description": "A verified unrelated creative project.",
                "technologies": [],
                "facts": ["Built a verified creative project."],
            },
            {
                "name": "AI Resume Tailoring Platform",
                "description": "FastAPI platform automation with Python and local LLM integration.",
                "technologies": ["Python", "Linux"],
                "facts": ["Built a verified FastAPI platform automation workflow."],
            },
        ]
    )
    profile = CandidateProfile.model_validate(payload)
    result = FactValidator(profile).validate(
        TailoredResume(
            headline=platform_job.role,
            professional_summary="",
            core_skills=[],
            experience=[],
        ),
        platform_job,
    )
    names = [project.name for project in result.resume.projects]
    assert len(names) == 2
    assert names[0] == "AI Resume Tailoring Platform"
    assert "Linux VPS" in names
