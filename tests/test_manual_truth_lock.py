import json

from docx import Document
from pypdf import PdfReader

from app.main import app
from app.services.manual_resume import manual_edit
from app.services.resume_editor import apply_resume_edits
from tests.test_safe_resume_edits import form_for
from tests.test_safe_resume_edits import saved_cv as saved_cv

IAM_SKILLS = [
    'Azure AD', 'Microsoft Identity & Access Management', 'AAD Group Management',
    'Tenant & Domain Management', 'Azure RBAC (Role-Based Access Control)',
    'Microsoft Graph API', 'Azure AD Roles and Administrators UI',
    'Privileged Identity Management (PIM)', 'Azure AD Access Reviews',
]


def switch(client, slug, enabled, **kwargs):
    return client.post(f'/preview/{slug}/truth-lock', data={
        'csrf_token': client.cookies.get('csrf_token'), 'enabled': enabled, **kwargs,
    }, follow_redirects=False)


def test_default_lock_rejects_new_skills_and_forged_edit_flag(saved_cv):
    client, slug, folder, original, _ = saved_cv
    result = client.post(f'/edit/{slug}', data=form_for(client, original,
        core_skills='\n'.join(IAM_SKILLS), truth_lock_enabled='false'), follow_redirects=False)
    assert result.status_code == 422
    assert app.state.storage.load_application(slug)[1] == original
    assert not (folder/'resume.pdf').exists()


def test_manual_mode_saves_plain_iam_names_in_real_exports(saved_cv):
    client, slug, folder, original, _ = saved_cv
    assert switch(client, slug, 'off').status_code == 303
    page = client.get(f'/edit/{slug}')
    assert 'OFF — manual CV' in page.text and 'Apply validation level' in page.text
    result = client.post(f'/edit/{slug}', data=form_for(client, original,
        core_skills='\n'.join(IAM_SKILLS)), follow_redirects=False)
    assert result.status_code == 303, result.text
    _, resume, generated, metadata = app.state.storage.load_application(slug)
    assert resume.core_skills == IAM_SKILLS
    assert resume.selected_skill_ids == []
    assert generated == original
    assert metadata['truth_lock_enabled'] is False
    assert 'Manual' in metadata['candidate_data_status']
    pdf_text = '\n'.join(page.extract_text() for page in PdfReader(folder/'resume.pdf').pages)
    document = Document(folder/'resume.docx')
    doc_text = '\n'.join(p.text for p in document.paragraphs)
    doc_text += '\n' + '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    pdf_text = ' '.join(pdf_text.split())
    doc_text = ' '.join(doc_text.split())
    for skill in IAM_SKILLS:
        assert skill in pdf_text and skill in doc_text
        assert 'Basic ' + skill not in doc_text
    assert 'Truth Lock is off' in ' '.join(metadata['truth_lock_warnings'])
    assert switch(client, slug, 'on').status_code == 409
    assert app.state.storage.load_application(slug)[3]['truth_lock_enabled'] is False


def test_toggle_requires_csrf_and_known_mode(saved_cv):
    client, slug, _, _, _ = saved_cv
    assert switch(client, slug, 'off', csrf_token='invalid').status_code == 403
    assert switch(client, slug, 'anything').status_code == 422
    assert app.state.storage.load_application(slug)[3].get('truth_lock_enabled', True) is True
    assert switch(client, 'missing-resume', 'off').status_code == 404


def test_legacy_missing_metadata_defaults_to_on(saved_cv):
    client, slug, folder, _, _ = saved_cv
    data = json.loads((folder/'metadata.json').read_text())
    data.pop('truth_lock_enabled', None)
    (folder/'metadata.json').write_text(json.dumps(data))
    assert 'Truth Lock: ON' in client.get(f'/edit/{slug}').text


def test_changed_manual_claims_do_not_retain_verified_source_ids(profile, resume_payload):
    from app.models.resume import TailoredResume
    original = TailoredResume.model_validate(resume_payload)
    edited = apply_resume_edits(original, {
        'professional_summary': 'Manually supplied summary.', 'core_skills': 'Microsoft Graph API',
        'experience_0_bullet_0': 'Manually supplied bullet.',
    })
    result = manual_edit(original, edited, 'example')
    assert result.resume.summary_source_fact_ids == []
    assert result.resume.summary_attestation == ''
    assert result.resume.selected_skill_ids == []
    assert result.resume.experience[0].bullets[0].source_fact_ids == ['manual:example:experience:0:0']
    assert original.experience[0].bullets[0].source_fact_ids != result.resume.experience[0].bullets[0].source_fact_ids


def test_manual_mode_is_scoped_to_one_document(saved_cv):
    client, slug, _, original, job = saved_cv
    assert switch(client, slug, 'off').status_code == 303
    from app.models.job import JobAnalysis
    analysis = JobAnalysis(company=job.company, role=job.role, strong_matches=[], partial_matches=[],
                           missing_requirements=[], recommendation='REVIEW', match_level='LOW')
    other, _ = app.state.storage.save_application(job, analysis, original, [])
    assert app.state.storage.load_application(other)[3]['truth_lock_enabled'] is True
    result = client.post(f'/edit/{other}', data=form_for(client, original,
        core_skills='Azure AD'), follow_redirects=False)
    assert result.status_code == 422


def test_source_backed_document_can_enable_lock_again(saved_cv):
    client, slug, _, original, job = saved_cv
    from app.services.fact_validator import FactValidator
    source_backed = FactValidator(app.state.profile, app.state.skills_bank).validate(original, job).resume
    app.state.storage.save_current_resume(slug, source_backed, [])
    assert switch(client, slug, 'off').status_code == 303
    response = switch(client, slug, 'on')
    assert response.status_code == 303, response.text
    assert app.state.storage.load_application(slug)[3]['truth_lock_enabled'] is True


def test_partial_mode_exports_manual_skills_but_rejects_employment_claims(saved_cv):
    client, slug, folder, original, _ = saved_cv
    response = client.post(f'/preview/{slug}/truth-lock', data={
        'csrf_token': client.cookies.get('csrf_token'), 'level': '0.5'}, follow_redirects=False)
    assert response.status_code == 303
    metadata = app.state.storage.load_application(slug)[3]
    assert metadata['truth_lock_level'] == 0.5
    assert metadata['truth_lock_enabled'] is True
    page = client.get(f'/edit/{slug}')
    assert 'PARTIAL — manual Skills' in page.text
    assert 'type="range"' in page.text and 'step="0.1"' in page.text
    result = client.post(f'/edit/{slug}', data=form_for(client, original,
        core_skills='Microsoft Graph API',
        experience_0_bullet_0='Led Azure IAM deployments for ten years.'), follow_redirects=False)
    assert result.status_code == 422
    assert app.state.storage.load_application(slug)[1] == original
    result = client.post(f'/edit/{slug}', data=form_for(client, original,
        core_skills='Microsoft Graph API'), follow_redirects=False)
    assert result.status_code == 303, result.text
    saved = app.state.storage.load_application(slug)[1]
    assert saved.core_skills == ['Microsoft Graph API']
    assert saved.selected_skill_ids == []
    assert saved.experience == original.experience
    assert 'Microsoft Graph API' in '\n'.join(p.text for p in Document(folder/'resume.docx').paragraphs)
    assert switch(client, slug, 'on').status_code == 409
    assert app.state.storage.load_application(slug)[3]['truth_lock_level'] == 0.5
    # Export again via the existing presentation path: manual skills survive.
    result = client.post(f'/edit/{slug}', data=form_for(client, saved), follow_redirects=False)
    assert result.status_code == 303
    assert app.state.storage.load_application(slug)[1].core_skills == ['Microsoft Graph API']


def test_slider_rejects_invalid_levels_and_requires_csrf(saved_cv):
    client, slug, _, _, _ = saved_cv
    for level in ['-0.1', '1.1', 'nan', 'inf', 'arbitrary']:
        response = client.post(f'/preview/{slug}/truth-lock', data={
            'csrf_token': client.cookies.get('csrf_token'), 'level': level}, follow_redirects=False)
        assert response.status_code == 422
    response = client.post(f'/preview/{slug}/truth-lock', data={
        'csrf_token': 'wrong', 'level': '0.5'}, follow_redirects=False)
    assert response.status_code == 403
    assert app.state.storage.load_application(slug)[3]['truth_lock_enabled'] is True


def test_middle_levels_have_same_scoped_permissions(profile, resume_payload):
    from app.models.job import JobRequest
    from app.models.resume import TailoredResume
    from app.services.fact_validator import FactValidator
    from app.services.manual_resume import check_edit
    validator = FactValidator(profile)
    job = JobRequest(company='Example', role='Support', job_description='Python support and operational automation.')
    original = TailoredResume.model_validate(resume_payload)
    edited = apply_resume_edits(original, {'core_skills': 'Microsoft Graph API',
        'experience_0_bullet_0': 'Commercial Azure IAM specialist for ten years.'})
    for level in [0.1, 0.5, 0.9]:
        result = check_edit(validator, original, edited, job, 'one-cv', {'truth_lock_level': level})
        assert result.resume.core_skills == ['Microsoft Graph API']
        assert result.resume.selected_skill_ids == []
        assert any(issue['field'] == 'experience_0_bullet_0' for issue in result.issues)
        assert not any(issue['field'].startswith('Skills:') for issue in result.issues)
    full = check_edit(validator, original, edited, job, 'one-cv', {})
    assert any(issue['field'].startswith('Skills:') for issue in full.issues)
    manual = check_edit(validator, original, edited, job, 'one-cv', {'truth_lock_level': 0})
    assert not manual.issues
    assert manual.resume.experience[0].bullets[0].source_fact_ids[0].startswith('manual:')


def test_tailoring_is_chosen_before_analysis_and_does_not_change_saved_cv(saved_cv):
    client, slug, folder, original, job = saved_cv
    before = (folder/'resume.json').read_bytes()
    page = client.get(f'/edit/{slug}')
    assert 'Create a new version with different job tailoring' in page.text
    assert 'name="tailoring_level"' in page.text
    tasks = []
    for level in [0.2, 0.8]:
        response = client.post('/analyze', data={
            'csrf_token': client.cookies.get('csrf_token'), 'company': job.company,
            'role': job.role, 'job_description': job.job_description, 'tailoring_level': str(level),
        }, headers={'Accept': 'application/json'}, follow_redirects=False)
        assert response.status_code == 202, response.text
        task = app.state.task_manager.get_task(response.json()['task_id'])
        assert task.payload['job']['tailoring_level'] == level
        tasks.append(task)
    assert tasks[0].idempotency_key != tasks[1].idempotency_key
    assert (folder/'resume.json').read_bytes() == before
    assert app.state.storage.load_application(slug)[1] == original
