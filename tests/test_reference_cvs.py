import io
import json

import pytest
from docx import Document

from app.services.reference_cvs import ReferenceCVError, ReferenceCVLibrary


def docx_bytes(role: str = "Platform Engineer") -> bytes:
    document = Document()
    document.add_paragraph("Jane Example")
    document.add_paragraph(role)
    document.add_heading("Professional Summary", level=1)
    document.add_paragraph("Builds reliable Python platforms and automates operational workflows.")
    document.add_heading("Technical Skills", level=1)
    document.add_paragraph("Python, Linux, Docker, FastAPI")
    document.add_heading("Professional Experience", level=1)
    document.add_paragraph("• Automated repeatable support workflows with Python and clear documentation.")
    document.add_heading("Education", level=1)
    document.add_paragraph("Example University")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_reference_library_ingests_indexes_selects_and_removes_docx(tmp_path):
    library = ReferenceCVLibrary(tmp_path / "reference_cvs")
    record = library.upload("platform-engineer.docx", docx_bytes())

    assert record["status"] == "ready"
    assert record["target_role"] == "Platform Engineer"
    assert record["skills"] == ["Python", "Linux", "Docker", "FastAPI"]
    assert record["experience_bullets"]
    assert record["section_order"][:3] == ["summary", "skills", "experience"]
    assert (tmp_path / "reference_cvs" / "index.json").stat().st_mode & 0o777 == 0o600

    selected = library.select("Python Platform Engineer", "Build Docker services", limit=2)
    assert selected[0]["target_role"] == "Platform Engineer"
    assert "original_filename" not in selected[0]

    persisted = json.loads((tmp_path / "reference_cvs" / "index.json").read_text())
    assert persisted["items"][0]["original_filename"] == "platform-engineer.docx"

    library.remove(record["id"])
    assert library.list() == []
    assert not list((tmp_path / "reference_cvs").glob("*.docx"))


def test_reference_library_rejects_unsupported_and_oversized_files(tmp_path):
    library = ReferenceCVLibrary(tmp_path / "reference_cvs", max_bytes=20)
    with pytest.raises(ReferenceCVError, match="PDF and DOCX"):
        library.upload("resume.txt", b"plain text")
    with pytest.raises(ReferenceCVError, match="size limit"):
        library.upload("resume.pdf", b"x" * 21)


def test_layout_guide_is_derived_from_ingested_references(tmp_path):
    library = ReferenceCVLibrary(tmp_path / "reference_cvs")
    library.upload("platform.docx", docx_bytes())
    guide = library.layout_guide()

    assert 280 <= guide["summary_max_chars"] <= 700
    assert 8 <= guide["max_skills"] <= 16
    assert 2 <= guide["max_bullets_per_role"] <= 4
    assert guide["section_order"][0] == "summary"
